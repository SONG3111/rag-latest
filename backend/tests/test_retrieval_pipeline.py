"""Retrieval behaviours added on top of the recall/fuse/rerank core.

These run without an API key: the dense leg, the reranker, and the rewriter are all
injected as fakes, which is also what makes them deterministic.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field

import pytest
from docx import Document
from openpyxl import Workbook

from app.models import Workspace
from app.retrieval.bm25 import content_fingerprint
from app.retrieval.pipeline import (
    SCORE_SOURCE_FUSED,
    SCORE_SOURCE_RERANK,
    Retriever,
)
from app.retrieval.rewriter import QueryRewriter, RewriteResult
from app.services.files import index_file, save_upload


# --------------------------------------------------------------------------- #
# fakes
# --------------------------------------------------------------------------- #
class FakeEmbeddings:
    """Deterministic, dependency-free stand-in for a real embedding model."""

    def _vector(self, text: str) -> list[float]:
        buckets = [0.0] * 8
        for index, char in enumerate(text):
            buckets[index % 8] += ord(char) % 17
        return buckets

    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        return [self._vector(text) for text in texts]

    def embed_query(self, text: str) -> list[float]:
        return self._vector(text)


class FakeVectorStore:
    """Records the query and returns a canned ranking."""

    def __init__(self, ids: list[str]) -> None:
        self.ids = ids
        self.queries: list[str] = []

    def search(self, workspace_id, embeddings, query, top_k):
        self.queries.append(query)
        return self.ids[:top_k]


@dataclass
class ScriptedReranker:
    """Returns a fixed ordering with fixed scores."""

    order: list[int]
    scores: list[float]
    produces_relevance_scores: bool = True
    calls: list[str] = field(default_factory=list)

    def rerank(self, query, documents, top_n, fused_scores=None):
        self.calls.append(query)
        return self.order[:top_n], self.scores[:top_n]


class FixedRewriter:
    def __init__(self, rewritten: str, *, raises: bool = False) -> None:
        self.rewritten = rewritten
        self.raises = raises
        self.seen: list[tuple[str, list[str]]] = []

    def rewrite(self, query: str, history: list[str] | None = None) -> RewriteResult:
        self.seen.append((query, list(history or [])))
        if self.raises:
            return RewriteResult(
                query=query, rewritten=False, original=query, error="boom"
            )
        return RewriteResult(
            query=self.rewritten,
            rewritten=self.rewritten != query,
            original=query,
        )


# --------------------------------------------------------------------------- #
# fixtures
# --------------------------------------------------------------------------- #
def _seed(session, tmp_path, monkeypatch):
    from app.config import get_settings

    settings = get_settings()
    monkeypatch.setattr(settings, "data_dir", tmp_path / "data", raising=False)
    settings.ensure_directories()

    workspace = Workspace(name="管线测试")
    session.add(workspace)
    session.flush()

    book = Workbook()
    sheet = book.active
    sheet.title = "报销明细"
    sheet.append(["报销单号", "姓名", "金额"])
    sheet.append(["BX-001", "张伟", 3200])
    sheet.append(["BX-002", "李娜", 8600])
    buffer = io.BytesIO()
    book.save(buffer)
    record = save_upload(session, workspace, "明细.xlsx", buffer.getvalue())
    index_file(session, workspace.id, record)

    document = Document()
    document.add_heading("报销制度", level=1)
    document.add_paragraph("第二条 单笔报销金额不得超过 5000 元。")
    buffer = io.BytesIO()
    document.save(buffer)
    record = save_upload(session, workspace, "制度.docx", buffer.getvalue())
    index_file(session, workspace.id, record)

    session.flush()
    return workspace


def _child_ids(session, workspace_id: str) -> list[str]:
    from sqlalchemy import select

    from app.models import Chunk

    return list(
        session.scalars(
            select(Chunk.id).where(
                Chunk.workspace_id == workspace_id, Chunk.level == "child"
            )
        )
    )


# --------------------------------------------------------------------------- #
# deduplication
# --------------------------------------------------------------------------- #
def test_fingerprint_ignores_whitespace_and_trailing_context() -> None:
    assert content_fingerprint("单笔报销 5000 元。" * 40) == content_fingerprint(
        "单笔报销5000元。" * 40
    )
    assert content_fingerprint("甲") != content_fingerprint("乙")


def test_dedupe_keeps_first_occurrence_and_preserves_order() -> None:
    from app.models import Chunk

    chunks = {
        "a": Chunk(id="a", workspace_id="w", file_id="f", text="同一段内容"),
        "b": Chunk(id="b", workspace_id="w", file_id="f", text="同一段内容"),
        "c": Chunk(id="c", workspace_id="w", file_id="f", text="另一段内容"),
    }
    deduped = Retriever._dedupe(["a", "b", "c"], chunks)
    assert deduped == ["a", "c"]


def test_dedupe_drops_unknown_ids() -> None:
    from app.models import Chunk

    chunks = {"a": Chunk(id="a", workspace_id="w", file_id="f", text="内容")}
    assert Retriever._dedupe(["missing", "a"], chunks) == ["a"]


# --------------------------------------------------------------------------- #
# parent context
# --------------------------------------------------------------------------- #
def test_children_are_ranked_and_parents_supply_context(
    temp_session, tmp_path, monkeypatch
) -> None:
    workspace = _seed(temp_session, tmp_path, monkeypatch)
    child_ids = _child_ids(temp_session, workspace.id)
    assert child_ids

    retriever = Retriever(
        temp_session,
        workspace.id,
        embeddings=FakeEmbeddings(),
        vector_store=FakeVectorStore(child_ids),
        reranker=ScriptedReranker(
            order=list(range(len(child_ids))),
            scores=[0.9] * len(child_ids),
        ),
        rewriter=FixedRewriter("张伟的报销金额"),
    )
    hits = retriever.search("张伟报销多少", use_dense=True, use_rerank=True)
    assert hits

    hit = hits[0]
    # The citation points at the row; the context handed to the model is the bucket.
    assert "第" in hit.location and "行" in hit.location
    assert hit.parent_text is not None
    assert hit.parent_location is not None
    assert len(hit.parent_text) >= len(hit.text)
    assert hit.score_source == SCORE_SOURCE_RERANK


# --------------------------------------------------------------------------- #
# relevance gate
# --------------------------------------------------------------------------- #
def test_low_scores_are_gated_out_entirely(temp_session, tmp_path, monkeypatch) -> None:
    workspace = _seed(temp_session, tmp_path, monkeypatch)
    child_ids = _child_ids(temp_session, workspace.id)
    from app.config import get_settings

    default_threshold = get_settings().rerank_score_threshold

    retriever = Retriever(
        temp_session,
        workspace.id,
        embeddings=FakeEmbeddings(),
        vector_store=FakeVectorStore(child_ids),
        reranker=ScriptedReranker(
            order=list(range(len(child_ids))),
            # Every candidate sits below the configured floor, which is what a question
            # outside the corpus looks like once terse child text is being scored.
            scores=[default_threshold / 10] * len(child_ids),
        ),
        rewriter=FixedRewriter("公司年假怎么申请"),
    )
    hits = retriever.search("公司年假怎么申请")

    assert hits == []
    assert retriever.last_run.filtered_by_threshold > 0
    assert retriever.last_run.threshold_applied == pytest.approx(default_threshold)


def test_scores_above_the_floor_survive(temp_session, tmp_path, monkeypatch) -> None:
    workspace = _seed(temp_session, tmp_path, monkeypatch)
    child_ids = _child_ids(temp_session, workspace.id)
    retriever = Retriever(
        temp_session,
        workspace.id,
        embeddings=FakeEmbeddings(),
        vector_store=FakeVectorStore(child_ids),
        reranker=ScriptedReranker(
            order=list(range(len(child_ids))),
            scores=[0.92] + [0.01] * (len(child_ids) - 1),
        ),
        rewriter=FixedRewriter("报销上限"),
    )
    hits = retriever.search("报销上限")
    assert len(hits) == 1
    assert hits[0].score == pytest.approx(0.92)


def test_threshold_zero_disables_the_gate(temp_session, tmp_path, monkeypatch) -> None:
    workspace = _seed(temp_session, tmp_path, monkeypatch)
    child_ids = _child_ids(temp_session, workspace.id)
    retriever = Retriever(
        temp_session,
        workspace.id,
        embeddings=FakeEmbeddings(),
        vector_store=FakeVectorStore(child_ids),
        reranker=ScriptedReranker(
            order=list(range(len(child_ids))),
            scores=[0.01] * len(child_ids),
        ),
        rewriter=FixedRewriter("任意"),
    )
    hits = retriever.search("任意", relevance_threshold=0.0)
    assert hits
    assert retriever.last_run.threshold_applied is None


def test_gate_is_skipped_when_scores_are_not_relevance(
    temp_session, tmp_path, monkeypatch
) -> None:
    """A noop reranker returns fused ranks, which a relevance floor cannot filter."""
    workspace = _seed(temp_session, tmp_path, monkeypatch)
    child_ids = _child_ids(temp_session, workspace.id)
    retriever = Retriever(
        temp_session,
        workspace.id,
        embeddings=FakeEmbeddings(),
        vector_store=FakeVectorStore(child_ids),
        reranker=ScriptedReranker(
            order=list(range(len(child_ids))),
            scores=[0.001] * len(child_ids),
            produces_relevance_scores=False,
        ),
        rewriter=FixedRewriter("任意"),
    )
    hits = retriever.search("任意")
    assert hits
    assert retriever.last_run.threshold_applied is None
    assert hits[0].score_source == SCORE_SOURCE_FUSED


# --------------------------------------------------------------------------- #
# rewriting
# --------------------------------------------------------------------------- #
def test_rewritten_query_is_used_for_both_legs(
    temp_session, tmp_path, monkeypatch
) -> None:
    workspace = _seed(temp_session, tmp_path, monkeypatch)
    child_ids = _child_ids(temp_session, workspace.id)
    store = FakeVectorStore(child_ids)
    reranker = ScriptedReranker(
        order=list(range(len(child_ids))), scores=[0.9] * len(child_ids)
    )
    rewriter = FixedRewriter("报销金额上限是多少")

    retriever = Retriever(
        temp_session,
        workspace.id,
        embeddings=FakeEmbeddings(),
        vector_store=store,
        reranker=reranker,
        rewriter=rewriter,
    )
    retriever.search("报销咋整", history=["用户: 报销制度怎么规定"])

    assert store.queries == ["报销金额上限是多少"]
    assert reranker.calls == ["报销金额上限是多少"]
    assert retriever.last_run.rewritten is True
    assert retriever.last_run.original_query == "报销咋整"


def test_history_reaches_the_rewriter(temp_session, tmp_path, monkeypatch) -> None:
    workspace = _seed(temp_session, tmp_path, monkeypatch)
    child_ids = _child_ids(temp_session, workspace.id)
    rewriter = FixedRewriter("第三条的内容是什么")
    retriever = Retriever(
        temp_session,
        workspace.id,
        embeddings=FakeEmbeddings(),
        vector_store=FakeVectorStore(child_ids),
        reranker=ScriptedReranker(order=[0], scores=[0.9]),
        rewriter=rewriter,
    )
    retriever.search("那第三条呢", history=["用户: 第二条怎么规定", "助手: 第二条是..."] )
    assert rewriter.seen[0][0] == "那第三条呢"
    assert rewriter.seen[0][1] == ["用户: 第二条怎么规定", "助手: 第二条是..."]


def test_rewrite_is_skipped_when_disabled(temp_session, tmp_path, monkeypatch) -> None:
    workspace = _seed(temp_session, tmp_path, monkeypatch)
    child_ids = _child_ids(temp_session, workspace.id)
    store = FakeVectorStore(child_ids)
    rewriter = FixedRewriter("不该被用到")

    retriever = Retriever(
        temp_session,
        workspace.id,
        embeddings=FakeEmbeddings(),
        vector_store=store,
        reranker=ScriptedReranker(order=[0], scores=[0.9]),
        rewriter=rewriter,
    )
    retriever.search("原问题", use_rewrite=False)

    assert rewriter.seen == []
    assert store.queries == ["原问题"]


# --------------------------------------------------------------------------- #
# rewriter unit behaviour
# --------------------------------------------------------------------------- #
class _ScriptedLLM:
    def __init__(self, content) -> None:
        self.content = content
        self.calls = 0

    def invoke(self, messages):
        self.calls += 1
        if isinstance(self.content, Exception):
            raise self.content
        from langchain_core.messages import AIMessage

        return AIMessage(content=self.content)


def test_rewriter_strips_quotes_and_takes_first_line() -> None:
    rewriter = QueryRewriter(llm=_ScriptedLLM('"报销流程是什么"\n多余的解释'))
    result = rewriter.rewrite("报销咋整")
    assert result.query == "报销流程是什么"
    assert result.rewritten is True


def test_rewriter_falls_back_on_exception() -> None:
    rewriter = QueryRewriter(llm=_ScriptedLLM(RuntimeError("network down")))
    result = rewriter.rewrite("报销咋整")
    assert result.query == "报销咋整"
    assert result.rewritten is False
    assert result.error


def test_rewriter_rejects_empty_output() -> None:
    rewriter = QueryRewriter(llm=_ScriptedLLM("   "))
    result = rewriter.rewrite("报销咋整")
    assert result.query == "报销咋整"
    assert result.rewritten is False


def test_rewriter_rejects_runaway_output() -> None:
    rewriter = QueryRewriter(llm=_ScriptedLLM("很长" * 500))
    result = rewriter.rewrite("报销咋整")
    assert result.query == "报销咋整"
    assert result.rewritten is False


def test_rewriter_disabled_makes_no_model_call() -> None:
    from app.config import get_settings

    settings = get_settings()
    llm = _ScriptedLLM("不该被调用")
    rewriter = QueryRewriter(settings, llm=llm)
    original = settings.query_rewrite_enabled
    settings.query_rewrite_enabled = False
    try:
        result = rewriter.rewrite("报销咋整")
    finally:
        settings.query_rewrite_enabled = original
    assert result.query == "报销咋整"
    assert llm.calls == 0


# --------------------------------------------------------------------------- #
# write-lock discipline
# --------------------------------------------------------------------------- #
def test_index_file_embeds_before_any_db_write(temp_session, tmp_path, monkeypatch) -> None:
    """Reindexing must not hold SQLite's write lock while the model embeds.

    It used to flush an ``indexing`` status row first and embed inside that
    transaction, so a concurrent approval click queued behind the whole
    embedding run and died with ``database is locked``. Embedding must happen
    before the first write of the indexing burst.
    """
    import io as _io

    from sqlalchemy import event

    from app.config import get_settings
    from app.services.files import index_file, save_upload

    settings = get_settings()
    monkeypatch.setattr(settings, "data_dir", tmp_path / "data", raising=False)
    settings.ensure_directories()

    workspace = Workspace(name="锁窗口测试")
    temp_session.add(workspace)
    temp_session.flush()

    book = Workbook()
    book.active.append(["产品", "单价"])
    book.active.append(["甲", 10])
    buffer = _io.BytesIO()
    book.save(buffer)
    record = save_upload(temp_session, workspace, "小表.xlsx", buffer.getvalue())
    temp_session.flush()

    state = {"writes_started": False}

    def _mark_before_flush(*_args, **_kwargs) -> None:
        state["writes_started"] = True

    event.listen(temp_session, "before_flush", _mark_before_flush)

    class SpyEmbeddings:
        def embed_documents(self, texts):
            assert not state["writes_started"], (
                "embedding ran after the write burst began; the SQLite write "
                "lock would be held across the model call"
            )
            return [[float(len(text)) % 7] * 4 for text in texts]

        def embed_query(self, text):
            return [0.0] * 4

    upserts: list[int] = []

    class RecordingStore:
        def upsert_vectors(self, workspace_id, chunk_ids, vectors, payloads):
            upserts.append(len(chunk_ids))
            assert len(vectors) == len(chunk_ids)
            return len(chunk_ids)

    result = index_file(
        temp_session,
        workspace.id,
        record,
        embeddings=SpyEmbeddings(),
        vector_store=RecordingStore(),
    )

    assert result.chunk_count > 0
    assert result.vector_count == result.chunk_count
    assert upserts == [result.chunk_count]
    assert state["writes_started"]
