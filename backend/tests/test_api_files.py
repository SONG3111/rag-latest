"""File lifecycle tests over the REST surface: upload, download, delete, reindex.

The upload path's happy flow is covered in ``test_api.py``; these tests pin the
rest of the contract a file manager depends on: downloads round-trip bytes,
deletes remove the record, the disk copy, and the index together, reindex
replaces (never duplicates) chunks, and workspace deletion cleans the directory.
"""

from __future__ import annotations

import io

import pytest
from docx import Document
from httpx import ASGITransport, AsyncClient
from openpyxl import Workbook

pytestmark = pytest.mark.anyio

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def workbook_bytes(sheet_rows: int = 2) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "销售"
    sheet.append(["产品", "销售额"])
    for index in range(1, sheet_rows):
        sheet.append([f"产品{index}", index * 100])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def docx_bytes() -> bytes:
    document = Document()
    document.add_heading("报销制度", level=1)
    document.add_paragraph("第二条 单笔报销金额不得超过 5000 元。")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


async def _create_with_files(client: AsyncClient, *files: tuple[str, bytes]) -> str:
    workspace_id = (
        await client.post("/api/workspaces", json={"name": "文件管理"})
    ).json()["id"]
    uploaded = await client.post(
        f"/api/workspaces/{workspace_id}/files",
        files=[
            ("files", (name, content, XLSX_MIME if name.endswith("xlsx") else "application/msword"))
            for name, content in files
        ],
    )
    assert uploaded.status_code == 201
    return workspace_id


def _workspace_dir(app, workspace_id: str):
    from app.config import get_settings

    return get_settings().workspace_dir(workspace_id)


async def test_download_round_trips_the_original_bytes(app_with_temp_storage) -> None:
    transport = ASGITransport(app=app_with_temp_storage)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        async with app_with_temp_storage.router.lifespan_context(app_with_temp_storage):
            payload = workbook_bytes()
            workspace_id = await _create_with_files(client, ("销售表.xlsx", payload))
            file_id = (
                await client.get(f"/api/workspaces/{workspace_id}/files")
            ).json()[0]["id"]

            downloaded = await client.get(
                f"/api/workspaces/{workspace_id}/files/{file_id}/download"
            )
            assert downloaded.status_code == 200
            assert downloaded.content == payload
            # Non-ASCII filenames arrive percent-encoded in the header.
            assert "attachment" in downloaded.headers.get("content-disposition", "")

            missing_id = "0" * 32
            unknown = await client.get(
                f"/api/workspaces/{workspace_id}/files/{missing_id}/download"
            )
            assert unknown.status_code == 404


async def test_download_after_disk_loss_is_404(app_with_temp_storage) -> None:
    transport = ASGITransport(app=app_with_temp_storage)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        async with app_with_temp_storage.router.lifespan_context(app_with_temp_storage):
            workspace_id = await _create_with_files(client, ("销售表.xlsx", workbook_bytes()))
            record = (await client.get(f"/api/workspaces/{workspace_id}/files")).json()[0]

            (_workspace_dir(app_with_temp_storage, workspace_id) / record["rel_path"]).unlink()

            response = await client.get(
                f"/api/workspaces/{workspace_id}/files/{record['id']}/download"
            )
            assert response.status_code == 404
            assert "missing" in response.json()["detail"]


async def test_delete_file_removes_record_disk_and_index_entries(
    app_with_temp_storage,
) -> None:
    transport = ASGITransport(app=app_with_temp_storage)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        async with app_with_temp_storage.router.lifespan_context(app_with_temp_storage):
            workspace_id = await _create_with_files(client, ("销售表.xlsx", workbook_bytes()))
            record = (await client.get(f"/api/workspaces/{workspace_id}/files")).json()[0]

            deleted = await client.delete(
                f"/api/workspaces/{workspace_id}/files/{record['id']}"
            )
            assert deleted.status_code == 204

            listing = await client.get(f"/api/workspaces/{workspace_id}/files")
            assert listing.json() == []
            assert not (
                _workspace_dir(app_with_temp_storage, workspace_id) / record["rel_path"]
            ).exists()

            # The chunk rows were deleted with the file, so keyword retrieval is empty.
            from sqlalchemy import select

            from app.models import Chunk

            session = app_with_temp_storage.state.test_session_factory()
            try:
                remaining = session.scalars(
                    select(Chunk).where(Chunk.workspace_id == workspace_id)
                ).all()
                assert remaining == []
            finally:
                session.close()

            again = await client.delete(
                f"/api/workspaces/{workspace_id}/files/{record['id']}"
            )
            assert again.status_code == 404


async def test_reindex_endpoints_replace_not_duplicate(app_with_temp_storage) -> None:
    transport = ASGITransport(app=app_with_temp_storage)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        async with app_with_temp_storage.router.lifespan_context(app_with_temp_storage):
            workspace_id = await _create_with_files(
                client,
                ("销售表.xlsx", workbook_bytes()),
                ("制度.docx", docx_bytes()),
            )

            record = (await client.get(f"/api/workspaces/{workspace_id}/files")).json()[0]
            first = await client.post(
                f"/api/workspaces/{workspace_id}/files/{record['id']}/reindex"
            )
            assert first.status_code == 200
            assert first.json()["status"] == "indexed"
            assert first.json()["chunk_count"] > 0

            whole = await client.post(f"/api/workspaces/{workspace_id}/reindex")
            assert whole.status_code == 200
            results = whole.json()
            assert len(results) == 2
            assert all(item["status"] == "indexed" for item in results)
            assert all(item["chunk_count"] > 0 for item in results)

            # Re-indexing replaced the derived rows: no file contributes twice.
            from sqlalchemy import func, select

            from app.models import Chunk

            session = app_with_temp_storage.state.test_session_factory()
            try:
                counts = session.execute(
                    select(Chunk.file_id, func.count(Chunk.id))
                    .where(Chunk.workspace_id == workspace_id, Chunk.level == "child")
                    .group_by(Chunk.file_id)
                ).all()
                per_file = dict(counts)
                for item in results:
                    assert per_file[item["file_id"]] == item["chunk_count"]
            finally:
                session.close()


async def test_uploading_the_same_name_twice_never_clobbers(app_with_temp_storage) -> None:
    transport = ASGITransport(app=app_with_temp_storage)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        async with app_with_temp_storage.router.lifespan_context(app_with_temp_storage):
            first = workbook_bytes(sheet_rows=2)
            second = workbook_bytes(sheet_rows=5)
            workspace_id = await _create_with_files(client, ("销售表.xlsx", first))
            again = await client.post(
                f"/api/workspaces/{workspace_id}/files",
                files=[("files", ("销售表.xlsx", second, XLSX_MIME))],
            )
            assert again.status_code == 201
            assert again.json()[0]["rel_path"] == "销售表(1).xlsx"

            listing = (await client.get(f"/api/workspaces/{workspace_id}/files")).json()
            assert {item["rel_path"] for item in listing} == {
                "销售表.xlsx",
                "销售表(1).xlsx",
            }
            on_disk = {
                path.name
                for path in _workspace_dir(app_with_temp_storage, workspace_id).iterdir()
            }
            assert on_disk == {"销售表.xlsx", "销售表(1).xlsx"}


async def test_workspace_delete_removes_directory_and_everything_in_it(
    app_with_temp_storage,
) -> None:
    transport = ASGITransport(app=app_with_temp_storage)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        async with app_with_temp_storage.router.lifespan_context(app_with_temp_storage):
            workspace_id = await _create_with_files(client, ("销售表.xlsx", workbook_bytes()))
            directory = _workspace_dir(app_with_temp_storage, workspace_id)
            assert directory.exists()

            deleted = await client.delete(f"/api/workspaces/{workspace_id}")
            assert deleted.status_code == 204
            assert not directory.exists()
            assert (await client.get(f"/api/workspaces/{workspace_id}")).status_code == 404
            assert (await client.get("/api/workspaces")).json() == []
