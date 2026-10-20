"""Multi-agent delegation: the workspace agent hands analysis to a sub-agent.

The control flow pinned here is deterministic and driven by a scripted model, so
no real LLM is contacted (AGENTS.md). What matters is *who runs what*: a single
``analyze_and_report`` call runs the read-only sub-graph, the sub-graph's tools
are real MCP tools, and the parent relays the report it gets back.
"""

from __future__ import annotations

import asyncio
import json
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from langchain_core.messages import AIMessage, AIMessageChunk
from openpyxl import Workbook

from app.agent.analysis import AnalysisResult, DataAnalysisAgent
from app.agent.graph import (
    ANALYSIS_TOOL_NAME,
    KNOWLEDGE_TOOL_NAME,
    WorkspaceAgent,
)
from app.mcp_client import McpOfficeClient
from app.models import Workspace

REPORT = (
    "## 数据概览\n"
    "数据源：销售表.xlsx 的「销售」工作表，共 3 行。\n\n"
    "## 关键指标\n\n"
    "| 指标 | 数值 |\n|---|---|\n| 销售额 合计 | 6000 |\n\n"
    "## 数据出处\n[销售表.xlsx · 销售!第2行]\n"
)

WRITE_TOOL_NAMES = {
    "update_cells", "set_formula", "insert_rows", "delete_rows", "format_range",
    "insert_columns", "delete_columns", "copy_range", "delete_range", "merge_cells",
    "unmerge_cells", "find_replace", "manage_sheets", "replace_text",
    "update_table_cell",
}


class ScriptedLLM:
    """Same double as test_agent.py's: yields each scripted message, then stops.

    Shared between the parent and the sub-agent in these tests (the parent passes
    its own llm down), so the script is consumed in strict call order — which is
    exactly what makes the delegation sequence observable.
    """

    def __init__(self, script: list[AIMessage]) -> None:
        self.script = list(script)
        self.calls: list[list[Any]] = []

    def bind_tools(self, tools: list[Any]) -> "ScriptedLLM":
        self.bound_tools = tools
        return self

    def invoke(self, messages: list[Any]) -> AIMessage:
        self.calls.append(list(messages))
        return self.script.pop(0) if self.script else AIMessage(content="结束")

    async def astream(self, messages: list[Any]):
        self.calls.append(list(messages))
        message = self.script.pop(0) if self.script else AIMessage(content="结束")
        yield AIMessageChunk(
            content=message.content,
            tool_calls=list(getattr(message, "tool_calls", None) or []),
        )


@pytest.fixture()
def analysis_env(temp_session, tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    from app.config import get_settings

    settings = get_settings()
    monkeypatch.setattr(settings, "data_dir", tmp_path / "data", raising=False)
    settings.ensure_directories()

    workspace = Workspace(name="分析测试")
    temp_session.add(workspace)
    temp_session.flush()

    directory = settings.workspace_dir(workspace.id)
    directory.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    sheet = book.active
    sheet.title = "销售"
    sheet.append(["产品", "区域", "销售额"])
    sheet.append(["A型", "华东", 1000])
    sheet.append(["B型", "华北", 2000])
    sheet.append(["C型", "华南", 3000])
    book.save(directory / "销售表.xlsx")

    # A Word document too: the statistics tools must reject it as a type mismatch.
    document = Document()
    document.add_paragraph("第一条 员工出差需提前申请。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "限额"
    table.cell(1, 1).text = "5000"
    document.save(str(directory / "制度.docx"))

    return workspace, directory


def _call(name: str, args: dict, call_id: str = "call-1") -> AIMessage:
    return AIMessage(content="", tool_calls=[{"name": name, "args": args, "id": call_id}])


async def _run_agent(agent: WorkspaceAgent, message: str) -> list:
    events = []
    async for event in agent.astream(message):
        events.append(event)
    return events


def _drive(workspace, session, scripted: ScriptedLLM, message: str) -> list:
    """Run one turn against a real MCP subprocess, then tear the client down."""

    async def scenario():
        from app.config import get_settings

        client = McpOfficeClient(workspace_root=get_settings().workspaces_dir)
        await client.start()
        try:
            agent = WorkspaceAgent(session, workspace, client, llm=scripted)
            return await _run_agent(agent, message)
        finally:
            await client.stop()

    return asyncio.run(scenario())


def test_one_delegation_runs_the_sub_agent_and_relays_its_report(
    analysis_env, temp_session
) -> None:
    workspace, _ = analysis_env
    scripted = ScriptedLLM(
        [
            # main agent: decide to delegate
            _call(
                ANALYSIS_TOOL_NAME,
                {"request": "按区域汇总销售额，生成一份报告"},
                "parent-1",
            ),
            # sub-agent: actually take the measurement
            _call(
                "describe_column",
                {"path": "销售表.xlsx", "sheet_name": "销售", "column": "销售额"},
                "sub-1",
            ),
            # sub-agent: write the report (no tool calls -> loop ends)
            AIMessage(content=REPORT),
            # main agent: relay it
            AIMessage(content=REPORT),
        ]
    )

    events = _drive(workspace, temp_session, scripted, "帮我分析销售表并生成报告")
    kinds = [event.type for event in events]

    assert "error" not in kinds
    assert "tool_call" in kinds
    delegated = next(event.data for event in events if event.type == "tool_call")
    assert delegated["tool"] == ANALYSIS_TOOL_NAME
    assert "分析" in delegated["label"]

    tool_result = next(event.data for event in events if event.type == "tool_result")
    payload = json.loads(tool_result["content"])
    assert payload["ok"] is True
    assert payload["data"]["report"] == REPORT.strip()
    assert payload["data"]["steps"] == [
        {"tool": "describe_column", "args": {"path": "销售表.xlsx", "sheet_name": "销售", "column": "销售额"}}
    ]

    # The report is the answer: the parent relays it instead of inventing one.
    done = next(event.data for event in events if event.type == "done")
    assert REPORT.strip() in done["content"]

    # Exactly one delegation: two model calls for the parent, two for the sub-agent.
    # A higher count would mean the loop re-entered instead of stopping.
    assert len(scripted.calls) == 4

    # The sub-agent's statistic really went through the MCP subprocess.
    stat_messages = [
        message
        for call in scripted.calls
        for message in call
        if getattr(message, "name", None) == "describe_column"
    ]
    assert stat_messages, "the statistics tool should have executed"
    assert json.loads(stat_messages[0].content)["data"]["sum"] == 6000


def test_sub_agent_tools_are_read_only(analysis_env, temp_session) -> None:
    """A missing write tool is the guarantee, not the prompt asking nicely."""
    workspace, _ = analysis_env

    async def scenario():
        from app.config import get_settings

        client = McpOfficeClient(workspace_root=get_settings().workspaces_dir)
        await client.start()
        try:
            sub_agent = DataAnalysisAgent(temp_session, workspace, client)
            names = sorted(tool.name for tool in sub_agent.build_tools())
            read_only = sorted(
                name for name, profile in client.profiles().items() if profile.read_only
            )
            return names, read_only
        finally:
            await client.stop()

    names, read_only = asyncio.run(scenario())

    assert names == sorted([KNOWLEDGE_TOOL_NAME, *read_only])
    assert WRITE_TOOL_NAMES.isdisjoint(names)
    for expected in ("describe_column", "aggregate", "sort_by", "read_range", "calculate"):
        assert expected in names


def test_report_citations_are_surfaced_on_the_parent_turn(
    analysis_env, temp_session, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Citations the sub-agent collected must reach the parent's citation panel."""
    workspace, _ = analysis_env
    citation = {
        "file": "销售表.xlsx",
        "location": "销售!第2行",
        "snippet": "A型 华东 1000",
        "score": 0.9,
        "score_source": "rerank",
    }

    class StubAnalysisAgent:
        def __init__(self, *_args, **_kwargs) -> None:
            pass

        async def arun(self, _request: str) -> AnalysisResult:
            return AnalysisResult(
                report=REPORT,
                citations=[citation],
                steps=[{"tool": "describe_column", "args": {}}],
            )

    import app.agent.analysis as analysis_module

    monkeypatch.setattr(analysis_module, "DataAnalysisAgent", StubAnalysisAgent)

    scripted = ScriptedLLM(
        [
            _call(ANALYSIS_TOOL_NAME, {"request": "分析"}, "parent-1"),
            AIMessage(content=REPORT),
        ]
    )
    events = _drive(workspace, temp_session, scripted, "分析一下并生成报告")

    citation_events = [event.data for event in events if event.type == "citations"]
    assert citation_events, "the report's citations should reach the stream"
    assert citation_events[0]["items"][0]["file"] == "销售表.xlsx"


def test_string_focus_files_still_delegates(analysis_env, temp_session) -> None:
    """Regression: the model sometimes emits ``focus_files`` as a string.

    Live testing had that shape rejected by the tool schema, which cost two
    identical retries before the delegation could start; the parameter now accepts
    both shapes and splits a delimited string.
    """
    workspace, _ = analysis_env
    scripted = ScriptedLLM(
        [
            _call(
                ANALYSIS_TOOL_NAME,
                {"request": "分析销售数据", "focus_files": "销售表.xlsx"},
                "parent-1",
            ),
            AIMessage(content=REPORT),
            AIMessage(content=REPORT),
        ]
    )

    events = _drive(workspace, temp_session, scripted, "分析一下销售表")

    assert "error" not in [event.type for event in events]
    tool_result = next(event.data for event in events if event.type == "tool_result")
    payload = json.loads(tool_result["content"])
    assert payload["ok"] is True

    # The hint reached the sub-agent as a rendered prompt, not as a rejected arg.
    sub_prompt = str(scripted.calls[1][-1].content)
    assert "只分析这些文件：销售表.xlsx" in sub_prompt


def test_word_document_gets_a_type_hint_not_a_corruption_claim(
    analysis_env, temp_session
) -> None:
    """A statistics call on a .docx must not read like "your file is damaged"."""
    workspace, _ = analysis_env
    scripted = ScriptedLLM(
        [
            _call(
                ANALYSIS_TOOL_NAME,
                {"request": "统计制度.docx 里的限额"},
                "parent-1",
            ),
            _call(
                "describe_column",
                {"path": "制度.docx", "sheet_name": "Sheet1", "column": "限额"},
                "sub-1",
            ),
            AIMessage(content=REPORT),
            AIMessage(content=REPORT),
        ]
    )

    events = _drive(workspace, temp_session, scripted, "分析一下制度.docx")

    stat_message = next(
        message
        for call in scripted.calls
        for message in call
        if getattr(message, "name", None) == "describe_column"
    )
    payload = json.loads(stat_message.content)
    assert payload["error"]["code"] == "unsupported_document"
    hint = payload["error"]["hint"]
    # The model is pointed at the Word tools instead of being told to re-upload.
    assert "Excel" in hint and "read_table" in hint

    # One rejected sub-tool call does not break the turn.
    delegated = next(event.data for event in events if event.type == "tool_result")
    assert json.loads(delegated["content"])["ok"] is True


def test_empty_report_fails_instead_of_letting_the_parent_invent_one(
    analysis_env, temp_session
) -> None:
    workspace, _ = analysis_env
    scripted = ScriptedLLM(
        [
            _call(ANALYSIS_TOOL_NAME, {"request": "分析"}, "parent-1"),
            # The sub-agent ends its turn with nothing visible.
            AIMessage(content=""),
            AIMessage(content="这次没能完成分析。"),
        ]
    )

    events = _drive(workspace, temp_session, scripted, "分析一下")

    tool_result = next(event.data for event in events if event.type == "tool_result")
    payload = json.loads(tool_result["content"])
    assert payload["ok"] is False
    assert payload["error"]["code"] == "analysis_failed"
    # The parent is told to say so rather than fabricate a report.
    assert payload["error"]["hint"]

    done = next(event.data for event in events if event.type == "done")
    assert done["content"] == "这次没能完成分析。"
