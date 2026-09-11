"""The approval loop is the safety-critical path, so it is tested end to end."""

from __future__ import annotations

import asyncio
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from app.mcp_client import McpOfficeClient
from app.models import OperationStatus, Workspace
from app.services.operations import (
    OperationError,
    apply_operation,
    build_tool_arguments,
    create_operation,
    extract_target_path,
    reject_operation,
    revert_operation,
    workspace_relative_path,
)


@pytest.fixture()
def workspace_with_file(temp_session, file_root: Path):
    from app.config import get_settings

    settings = get_settings()
    workspace = Workspace(name="审批测试")
    temp_session.add(workspace)
    temp_session.flush()

    directory = settings.workspace_dir(workspace.id)
    directory.mkdir(parents=True, exist_ok=True)
    target = directory / "销售表.xlsx"

    book = Workbook()
    sheet = book.active
    sheet.title = "销售"
    sheet.append(["产品", "销售额"])
    sheet.append(["A型", 1000])
    book.save(target)

    return workspace, target


def test_extract_target_path_accepts_known_keys() -> None:
    assert extract_target_path({"path": "a.xlsx"}) == "a.xlsx"
    assert extract_target_path({"filepath": "b.docx"}) == "b.docx"
    assert extract_target_path({"file_path": "c.xlsx"}) == "c.xlsx"


def test_extract_target_path_rejects_missing_path() -> None:
    with pytest.raises(OperationError):
        extract_target_path({"sheet_name": "Sheet1"})


def test_workspace_relative_path_normalizes_separators() -> None:
    assert workspace_relative_path("ws1", "sub\\file.xlsx") == "ws1/sub/file.xlsx"
    assert workspace_relative_path("ws1", "/leading.xlsx") == "ws1/leading.xlsx"


def test_build_tool_arguments_rewrites_only_path_keys() -> None:
    result = build_tool_arguments(
        "ws1", {"path": "a.xlsx", "sheet_name": "销售", "updates": [{"cell": "A1"}]}
    )
    assert result["path"] == "ws1/a.xlsx"
    assert result["sheet_name"] == "销售"
    assert result["updates"] == [{"cell": "A1"}]


async def _client() -> McpOfficeClient:
    from app.config import get_settings

    client = McpOfficeClient(workspace_root=get_settings().workspaces_dir)
    await client.start()
    return client


def test_unknown_tool_fails_closed() -> None:
    """A tool that is not in the server's inventory must be treated as gated."""
    client = McpOfficeClient()
    assert client.requires_approval("totally_made_up_tool") is True


def test_operations_gate_writes(workspace_with_file, temp_session) -> None:
    workspace, target = workspace_with_file
    original = target.read_bytes()

    async def scenario() -> None:
        client = await _client()
        try:
            proposal = create_operation(
                temp_session,
                workspace.id,
                "update_cells",
                {
                    "path": "销售表.xlsx",
                    "sheet_name": "销售",
                    "updates": [{"cell": "B2", "value": 1500}],
                },
                diff=[{"cell": "B2", "before": 1000, "after": 1500}],
            )
            temp_session.flush()

            # Creating a proposal must not touch the file.
            assert target.read_bytes() == original
            assert proposal.status is OperationStatus.proposed
            assert "B2" in proposal.summary

            result = await apply_operation(temp_session, proposal, client)
            assert result["changes"] == [{"cell": "B2", "before": 1000, "after": 1500}]
            assert proposal.status is OperationStatus.applied
            assert proposal.backup_path is not None
            assert Path(proposal.backup_path).exists()
            assert target.read_bytes() != original
            assert load_workbook(target)["销售"]["B2"].value == 1500

            # Reverting restores the original bytes.
            revert_operation(temp_session, proposal)
            assert target.read_bytes() == original
            assert load_workbook(target)["销售"]["B2"].value == 1000
        finally:
            await client.stop()

    asyncio.run(scenario())


def test_rejecting_leaves_file_byte_identical(workspace_with_file, temp_session) -> None:
    workspace, target = workspace_with_file
    original = target.read_bytes()

    proposal = create_operation(
        temp_session,
        workspace.id,
        "update_cells",
        {
            "path": "销售表.xlsx",
            "sheet_name": "销售",
            "updates": [{"cell": "B2", "value": 9999}],
        },
    )
    temp_session.flush()
    reject_operation(temp_session, proposal)

    assert proposal.status is OperationStatus.rejected
    assert target.read_bytes() == original


def test_applying_twice_is_rejected(workspace_with_file, temp_session) -> None:
    workspace, _ = workspace_with_file

    async def scenario() -> None:
        client = await _client()
        try:
            proposal = create_operation(
                temp_session,
                workspace.id,
                "update_cells",
                {
                    "path": "销售表.xlsx",
                    "sheet_name": "销售",
                    "updates": [{"cell": "B2", "value": 1}],
                },
            )
            temp_session.flush()
            await apply_operation(temp_session, proposal, client)
            with pytest.raises(OperationError):
                await apply_operation(temp_session, proposal, client)
        finally:
            await client.stop()

    asyncio.run(scenario())


def test_rejecting_twice_is_rejected(workspace_with_file, temp_session) -> None:
    workspace, _ = workspace_with_file
    proposal = create_operation(
        temp_session,
        workspace.id,
        "update_cells",
        {
            "path": "销售表.xlsx",
            "sheet_name": "销售",
            "updates": [{"cell": "B2", "value": 1}],
        },
    )
    temp_session.flush()
    reject_operation(temp_session, proposal)
    with pytest.raises(OperationError):
        reject_operation(temp_session, proposal)


def test_revert_requires_an_applied_operation(workspace_with_file, temp_session) -> None:
    workspace, _ = workspace_with_file
    proposal = create_operation(
        temp_session,
        workspace.id,
        "update_cells",
        {
            "path": "销售表.xlsx",
            "sheet_name": "销售",
            "updates": [{"cell": "B2", "value": 1}],
        },
    )
    temp_session.flush()
    with pytest.raises(OperationError):
        revert_operation(temp_session, proposal)


def test_summaries_are_human_readable() -> None:
    from app.services.operations import summarize_operation

    assert "C2" in summarize_operation(
        "update_cells",
        "表.xlsx",
        {"sheet_name": "销售", "updates": [{"cell": "C2"}]},
    )
    assert "替换" in summarize_operation(
        "replace_text", "制度.docx", {"find": "5000", "replace": "8000"}
    )
    assert "插入" in summarize_operation(
        "insert_rows", "表.xlsx", {"sheet_name": "销售", "start_row": 3, "count": 2}
    )
