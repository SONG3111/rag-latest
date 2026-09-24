"""POST /v1/chat/stream 的内部契约测试（backend-java 的调用视角）。

与 ``test_api_chat_stream.py`` 共享 ScriptedLLM 思路：本套件把同一套编排逻辑
对准无状态端点，钉住的差异点是——persist 帧替代 done、proposal 帧带全
arguments 且 operation_id 为 null、压缩书签随 persist 帧外传、以及历史消息
完全由请求携带（服务端不读库）。全程无真实模型调用。
"""

from __future__ import annotations

import io
import json
import re
from typing import Any

import pytest
from httpx import ASGITransport, AsyncClient
from langchain_core.messages import AIMessage, AIMessageChunk
from openpyxl import Workbook

from conftest import seed_workspace_file

pytestmark = pytest.mark.anyio


class ScriptedLLM:
    """Replays a fixed sequence of AI messages over the streaming surface."""

    def __init__(self, script: list[AIMessage]) -> None:
        self.script = list(script)
        self.calls: list[list[Any]] = []

    def bind_tools(self, tools: list[Any]) -> "ScriptedLLM":
        return self

    async def astream(self, messages: list[Any]):
        self.calls.append(list(messages))
        message = self.script.pop(0) if self.script else AIMessage(content="结束")
        yield AIMessageChunk(
            content=message.content,
            tool_calls=list(getattr(message, "tool_calls", None) or []),
        )


class ExplodingLLM(ScriptedLLM):
    async def astream(self, messages: list[Any]):
        self.calls.append(list(messages))
        raise RuntimeError("provider connection reset")
        yield  # pragma: no cover - makes this an async generator


def _call(name: str, args: dict, call_id: str = "call-1") -> AIMessage:
    return AIMessage(content="", tool_calls=[{"name": name, "args": args, "id": call_id}])


def parse_sse_frames(body: str) -> list[tuple[str, dict]]:
    frames: list[tuple[str, dict]] = []
    for block in re.split(r"\r?\n\r?\n", body):
        event = None
        data: list[str] = []
        for line in block.splitlines():
            if line.startswith("event:"):
                event = line.split(":", 1)[1].strip()
            elif line.startswith("data:"):
                data.append(line.split(":", 1)[1].strip())
        if event and data:
            frames.append((event, json.loads("\n".join(data))))
    return frames


def workbook_bytes() -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "销售"
    sheet.append(["产品", "销售额"])
    sheet.append(["A型", 1000])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


async def _upload_workbook(client: AsyncClient) -> str:
    """Stage the workbook the way Java's upload would (stateless service)."""
    workspace_id = "ws-chat-stream"
    seed_workspace_file(workspace_id, "销售表.xlsx", workbook_bytes())
    return workspace_id


async def _stream_turn(app, workspace_id: str, message: str, **extra) -> list[tuple[str, dict]]:
    payload = {
        "workspace_id": workspace_id,
        "run_id": "run-fixed-0001",
        "message": message,
        "files": ["销售表.xlsx"],
        **extra,
    }
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        response = await client.post("/v1/chat/stream", json=payload)
        assert response.status_code == 200
        return parse_sse_frames(response.text)


@pytest.fixture()
def scripted_model(monkeypatch: pytest.MonkeyPatch):
    def _install(script: list[AIMessage]) -> ScriptedLLM:
        from app.llm import providers

        model = ScriptedLLM(script)
        monkeypatch.setattr(
            providers, "build_resilient_chat_model", lambda settings: model
        )
        return model

    return _install


async def test_tool_turn_ends_with_persist_and_no_done(
    app_with_temp_storage, scripted_model
) -> None:
    transport = ASGITransport(app=app_with_temp_storage)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        async with app_with_temp_storage.router.lifespan_context(app_with_temp_storage):
            workspace_id = await _upload_workbook(client)
            scripted_model(
                [
                    _call(
                        "read_range",
                        {"path": "销售表.xlsx", "sheet_name": "销售", "start_cell": "A1"},
                    ),
                    AIMessage(content="表格里 A型 的销售额是 1000。"),
                ]
            )

            frames = await _stream_turn(app_with_temp_storage, workspace_id, "A型 卖了多少？")

    events = [name for name, _ in frames]
    assert "tool_call" in events and "tool_result" in events
    assert "done" not in events, "the stateless endpoint must not emit done"
    assert events[-1] == "persist", "persist is the terminating frame"

    persist = frames[-1][1]
    assert persist["status"] == "ok"
    assert persist["content"] == "表格里 A型 的销售额是 1000。"
    assert persist["proposals"] == []
    assert persist["trace_nodes"], "trace records ride the persist frame"
    assert persist["summary"] is None, "below the compaction watermark there is no summary"


async def test_history_comes_from_the_request_not_the_database(
    app_with_temp_storage, scripted_model
) -> None:
    transport = ASGITransport(app=app_with_temp_storage)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        async with app_with_temp_storage.router.lifespan_context(app_with_temp_storage):
            workspace_id = await _upload_workbook(client)
            scripted_model([AIMessage(content="接着上一轮说。")])

            frames = await _stream_turn(
                app_with_temp_storage,
                workspace_id,
                "继续",
                messages=[
                    {"role": "user", "content": "第一轮问题"},
                    {"role": "assistant", "content": "第一轮回答：上限 5000。"},
                ],
            )

    persist = frames[-1][1]
    assert persist["content"] == "接着上一轮说。"


async def test_write_turn_proposal_carries_arguments_and_null_operation_id(
    app_with_temp_storage, scripted_model
) -> None:
    transport = ASGITransport(app=app_with_temp_storage)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        async with app_with_temp_storage.router.lifespan_context(app_with_temp_storage):
            workspace_id = await _upload_workbook(client)
            from app.config import get_settings

            target = get_settings().workspace_dir(workspace_id) / "销售表.xlsx"
            original = target.read_bytes()

            updates = [{"cell": "B2", "value": 1500}]
            scripted_model(
                [
                    _call(
                        "update_cells",
                        {
                            "path": "销售表.xlsx",
                            "sheet_name": "销售",
                            "updates": updates,
                        },
                    ),
                    AIMessage(content="已生成修改提案，等待你确认。"),
                ]
            )

            frames = await _stream_turn(
                app_with_temp_storage, workspace_id, "把 B2 改成 1500"
            )

    proposal = next(payload for name, payload in frames if name == "proposal")
    assert proposal["tool"] == "update_cells"
    assert proposal["operation_id"] is None, "Java generates and back-fills the id"
    assert proposal["arguments"]["updates"] == updates
    assert proposal["diff"] == [{"cell": "B2", "before": 1000, "after": 1500}]

    # The gate held: nothing reached the file during the chat turn.
    assert target.read_bytes() == original

    persist = frames[-1][1]
    assert len(persist["proposals"]) == 1
    # The placeholder answer announces the pending proposal (compose_answer).
    assert "1 条待确认的修改提案" in persist["content"]

    # 无状态端点不落库：ai-service 在自己的存储树里没有任何关系状态
    # （app.db 归 backend-java 进程所有，且路径根本不在本服务的职责内）。
    from app.config import get_settings

    assert not (get_settings().data_dir / "app.db").exists()


async def test_provider_failure_yields_error_event_then_placeholder_persist(
    app_with_temp_storage, monkeypatch: pytest.MonkeyPatch
) -> None:
    from app.llm import providers

    monkeypatch.setattr(
        providers, "build_resilient_chat_model", lambda settings: ExplodingLLM([])
    )

    transport = ASGITransport(app=app_with_temp_storage)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        async with app_with_temp_storage.router.lifespan_context(app_with_temp_storage):
            workspace_id = await _upload_workbook(client)
            frames = await _stream_turn(app_with_temp_storage, workspace_id, "随便说点什么")

    events = [name for name, _ in frames]
    assert "error" in events, "a provider crash must surface as an SSE error"
    assert events[-1] == "persist"
    persist = frames[-1][1]
    assert persist["status"] == "ok"
    # The empty-answer placeholder (compose_answer) rides the persist frame.
    assert persist["content"].strip() != ""


async def test_without_mcp_the_endpoint_refuses_up_front(app_with_temp_storage) -> None:
    """没有可用 MCP 子进程时直接 503，消息不进编排（对齐旧版守卫语义）。"""
    transport = ASGITransport(app=app_with_temp_storage)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        response = await client.post(
            "/v1/chat/stream",
            json={"workspace_id": "ws", "run_id": "r", "message": "在吗"},
        )
        assert response.status_code == 503
