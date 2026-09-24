"""Retrieval behaviours added on top of the recall/fuse/rerank core.

These run without an API key: the dense leg, the reranker, and the rewriter are all
injected as fakes, which is also what makes them deterministic. The corpus is
built in memory (the shape backend-java's retrieval-corpus endpoint returns);
there is no database behind the retriever any more.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook

from app.retrieval.bm25 import content_fingerprint, token_counts
from app.retrieval.chunking import chunk_document_groups
from app.retrieval.pipeline import (
    SCORE_SOURCE_FUSED,
    SCORE_SOURCE_RERANK,
    ChunkRecord,
    Retriever,
)
from app.retrieval.rewriter import QueryRewriter, RewriteResult


# --------------------------------------------------------------------------- #
# fakes
# --------------------------------------------------------------------------- #
class FakeEmbeddings:
    """Deterministic, dependency-free stand-in for a real embedding model."""

    def _vector(self, text: str) -> list[float]:
        buckets = [0.0] * 8
        for index, char in enumerate(text):
            buckets[index % 8] += ord(char) % 17
        return buckets

    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        return [self._vector(text) for text in texts]

    def embed_query(self, text: str) -> list[float]:
        return self._vector(text)


class FakeVectorStore:
    """Records the query and returns a canned ranking."""

    def __init__(self, ids: list[str]) -> None:
        self.ids = ids
        self.queries: list[str] = []

    def search(self, workspace_id, embeddings, query, top_k):
        self.queries.append(query)
        return self.ids[:top_k]


@dataclass
class ScriptedReranker:
    """Returns a fixed ordering with fixed scores."""

    order: list[int]
    scores: list[float]
    produces_relevance_scores: bool = True
    calls: list[str] = field(default_factory=list)

    def rerank(self, query, documents, top_n, fused_scores=None):
        self.calls.append(query)
        return self.order[:top_n], self.scores[:top_n]


class FixedRewriter:
    def __init__(self, rewritten: str, *, raises: bool = False) -> None:
        self.rewritten = rewritten
        self.raises = raises
        self.seen: list[tuple[str, list[str]]] = []

    def rewrite(self, query: str, history: list[str] | None = None) -> RewriteResult:
        self.seen.append((query, list(history or [])))
        if self.raises:
            return RewriteResult(
                query=query, rewritten=False, original=query, error="boom"
            )
        return RewriteResult(
            query=self.rewritten,
            rewritten=self.rewritten != query,
            original=query,
        )


# --------------------------------------------------------------------------- #
# corpus
# --------------------------------------------------------------------------- #
FILE_IDS = {"明细.xlsx": "f-excel", "制度.docx": "f-word"}


def _corpus(tmp_path: Path) -> tuple[dict[str, str], list[ChunkRecord]]:
    """Chunk one Excel + one Word file the way /v1/index would, in memory."""
    book = Workbook()
    sheet = book.active
    sheet.title = "报销明细"
    sheet.append(["报销单号", "姓名", "金额"])
    sheet.append(["BX-001", "张伟", 3200])
    sheet.append(["BX-002", "李娜", 8600])
    excel_path = tmp_path / "明细.xlsx"
    book.save(excel_path)

    document = Document()
    document.add_heading("报销制度", level=1)
    document.add_paragraph("第二条 单笔报销金额不得超过 5000 元。")
    word_path = tmp_path / "制度.docx"
    document.save(str(word_path))

    records: list[ChunkRecord] = []
    for path, rel_path in ((excel_path, "明细.xlsx"), (word_path, "制度.docx")):
        file_id = FILE_IDS[rel_path]
        for group_index, group in enumerate(chunk_document_groups(path, rel_path)):
            parent_id = f"{file_id}-g{group_index}"
            parent_counts = token_counts(group.parent.text)
            records.append(
                ChunkRecord(
                    id=parent_id,
                    file_id=file_id,
                    parent_id=None,
                    level="parent",
                    text=group.parent.text,
                    location=group.parent.location,
                    meta=group.parent.meta,
                    token_counts=parent_counts,
                    token_length=sum(parent_counts.values()),
                )
            )
            for child_index, child in enumerate(group.children):
                child_counts = token_counts(child.text)
                records.append(
                    ChunkRecord(
                        id=f"{parent_id}-c{child_index}",
                        file_id=file_id,
                        parent_id=parent_id,
                        level="child",
                        text=child.text,
                        location=child.location,
                        meta=child.meta,
                        token_counts=child_counts,
                        token_length=sum(child_counts.values()),
                    )
                )
    # The corpus file map is {file_id: rel_path}, the retrieval-corpus shape.
    files = {file_id: rel_path for rel_path, file_id in FILE_IDS.items()}
    return files, records


def _child_ids(corpus) -> list[str]:
    _, records = corpus
    return [record.id for record in records if record.level == "child"]


# --------------------------------------------------------------------------- #
# deduplication
# --------------------------------------------------------------------------- #
def test_fingerprint_ignores_whitespace_and_trailing_context() -> None:
    assert content_fingerprint("单笔报销 5000 元。" * 40) == content_fingerprint(
        "单笔报销5000元。" * 40
    )
    assert content_fingerprint("甲") != content_fingerprint("乙")


def test_dedupe_keeps_first_occurrence_and_preserves_order() -> None:
    chunks = {
        "a": ChunkRecord(id="a", file_id="f", parent_id=None, level="child", text="同一段内容", location="l"),
        "b": ChunkRecord(id="b", file_id="f", parent_id=None, level="child", text="同一段内容", location="l"),
        "c": ChunkRecord(id="c", file_id="f", parent_id=None, level="child", text="另一段内容", location="l"),
    }
    deduped = Retriever._dedupe(["a", "b", "c"], chunks)
    assert deduped == ["a", "c"]


def test_dedupe_drops_unknown_ids() -> None:
    chunks = {
        "a": ChunkRecord(id="a", file_id="f", parent_id=None, level="child", text="内容", location="l")
    }
    assert Retriever._dedupe(["missing", "a"], chunks) == ["a"]


# --------------------------------------------------------------------------- #
# parent context
# --------------------------------------------------------------------------- #
def test_children_are_ranked_and_parents_supply_context(tmp_path) -> None:
    corpus = _corpus(tmp_path)
    child_ids = _child_ids(corpus)
    assert child_ids

    retriever = Retriever(
        "ws",
        corpus_loader=lambda: corpus,
        embeddings=FakeEmbeddings(),
        vector_store=FakeVectorStore(child_ids),
        reranker=ScriptedReranker(
            order=list(range(len(child_ids))),
            scores=[0.9] * len(child_ids),
        ),
        rewriter=FixedRewriter("张伟的报销金额"),
    )
    hits = retriever.search("张伟报销多少", use_dense=True, use_rerank=True)
    assert hits

    hit = hits[0]
    # The citation points at the row; the context handed to the model is the bucket.
    assert "第" in hit.location and "行" in hit.location
    assert hit.parent_text is not None
    assert hit.parent_location is not None
    assert len(hit.parent_text) >= len(hit.text)
    assert hit.score_source == SCORE_SOURCE_RERANK


# --------------------------------------------------------------------------- #
# corpus loader (M2: 语料经 /internal/retrieval-corpus 回读)
# --------------------------------------------------------------------------- #
def test_corpus_loader_is_read_once_per_search(tmp_path) -> None:
    """child 排序、parent 附上下文、file map 给出路径；一次 search 只回读一次
    （children/parents/file map 共享同一份 memoized 结果）。"""
    corpus = _corpus(tmp_path)
    child_ids = _child_ids(corpus)
    files, records = corpus

    loader_calls = []

    def loader():
        loader_calls.append(1)
        return files, records

    retriever = Retriever(
        "ws",
        corpus_loader=loader,
        embeddings=FakeEmbeddings(),
        vector_store=FakeVectorStore(child_ids),
        reranker=ScriptedReranker(
            order=list(range(len(child_ids))), scores=[0.9] * len(child_ids)
        ),
        rewriter=FixedRewriter("张伟的报销金额"),
    )
    hits = retriever.search("张伟报销多少", use_dense=True, use_rerank=True)
    assert hits
    assert len(loader_calls) == 1

    hit = hits[0]
    assert hit.rel_path in files.values()
    assert hit.parent_text is not None
    assert hit.parent_location is not None


def test_corpus_loader_misses_map_to_placeholder_file(tmp_path) -> None:
    """语料里 file_id 对不上文件表时的兜底：引用退化为「未知文件」。"""
    corpus = _corpus(tmp_path)
    child_ids = _child_ids(corpus)

    def loader():
        return {}, [
            ChunkRecord(
                id=child_ids[0],
                file_id="ghost-file",
                parent_id=None,
                level="child",
                text="张伟的报销金额 3200",
                location="报销明细!第2行",
            )
        ]

    retriever = Retriever(
        "ws",
        corpus_loader=loader,
        embeddings=FakeEmbeddings(),
        vector_store=FakeVectorStore(child_ids[:1]),
        rewriter=FixedRewriter("张伟"),
    )
    hits = retriever.search("张伟", use_dense=True, use_rerank=False)
    assert hits and hits[0].rel_path == "未知文件"


# --------------------------------------------------------------------------- #
# relevance gate
# --------------------------------------------------------------------------- #
def test_low_scores_are_gated_out_entirely(tmp_path) -> None:
    corpus = _corpus(tmp_path)
    child_ids = _child_ids(corpus)
    from app.config import get_settings

    default_threshold = get_settings().rerank_score_threshold

    retriever = Retriever(
        "ws",
        corpus_loader=lambda: corpus,
        embeddings=FakeEmbeddings(),
        vector_store=FakeVectorStore(child_ids),
        reranker=ScriptedReranker(
            order=list(range(len(child_ids))),
            # Every candidate sits below the configured floor, which is what a question
            # outside the corpus looks like once terse child text is being scored.
            scores=[default_threshold / 10] * len(child_ids),
        ),
        rewriter=FixedRewriter("公司年假怎么申请"),
    )
    hits = retriever.search("公司年假怎么申请")

    assert hits == []
    assert retriever.last_run.filtered_by_threshold > 0
    assert retriever.last_run.threshold_applied == pytest.approx(default_threshold)


def test_scores_above_the_floor_survive(tmp_path) -> None:
    corpus = _corpus(tmp_path)
    child_ids = _child_ids(corpus)
    retriever = Retriever(
        "ws",
        corpus_loader=lambda: corpus,
        embeddings=FakeEmbeddings(),
        vector_store=FakeVectorStore(child_ids),
        reranker=ScriptedReranker(
            order=list(range(len(child_ids))),
            scores=[0.92] + [0.01] * (len(child_ids) - 1),
        ),
        rewriter=FixedRewriter("报销上限"),
    )
    hits = retriever.search("报销上限")
    assert len(hits) == 1
    assert hits[0].score == pytest.approx(0.92)


def test_threshold_zero_disables_the_gate(tmp_path) -> None:
    corpus = _corpus(tmp_path)
    child_ids = _child_ids(corpus)
    retriever = Retriever(
        "ws",
        corpus_loader=lambda: corpus,
        embeddings=FakeEmbeddings(),
        vector_store=FakeVectorStore(child_ids),
        reranker=ScriptedReranker(
            order=list(range(len(child_ids))),
            scores=[0.01] * len(child_ids),
        ),
        rewriter=FixedRewriter("任意"),
    )
    hits = retriever.search("任意", relevance_threshold=0.0)
    assert hits
    assert retriever.last_run.threshold_applied is None


def test_gate_is_skipped_when_scores_are_not_relevance(tmp_path) -> None:
    """A noop reranker returns fused ranks, which a relevance floor cannot filter."""
    corpus = _corpus(tmp_path)
    child_ids = _child_ids(corpus)
    retriever = Retriever(
        "ws",
        corpus_loader=lambda: corpus,
        embeddings=FakeEmbeddings(),
        vector_store=FakeVectorStore(child_ids),
        reranker=ScriptedReranker(
            order=list(range(len(child_ids))),
            scores=[0.001] * len(child_ids),
            produces_relevance_scores=False,
        ),
        rewriter=FixedRewriter("任意"),
    )
    hits = retriever.search("任意")
    assert hits
    assert retriever.last_run.threshold_applied is None
    assert hits[0].score_source == SCORE_SOURCE_FUSED


# --------------------------------------------------------------------------- #
# rewriting
# --------------------------------------------------------------------------- #
def test_rewritten_query_is_used_for_both_legs(tmp_path) -> None:
    corpus = _corpus(tmp_path)
    child_ids = _child_ids(corpus)
    store = FakeVectorStore(child_ids)
    reranker = ScriptedReranker(
        order=list(range(len(child_ids))), scores=[0.9] * len(child_ids)
    )
    rewriter = FixedRewriter("报销金额上限是多少")

    retriever = Retriever(
        "ws",
        corpus_loader=lambda: corpus,
        embeddings=FakeEmbeddings(),
        vector_store=store,
        reranker=reranker,
        rewriter=rewriter,
    )
    retriever.search("报销咋整", history=["用户: 报销制度怎么规定"])

    assert store.queries == ["报销金额上限是多少"]
    assert reranker.calls == ["报销金额上限是多少"]
    assert retriever.last_run.rewritten is True
    assert retriever.last_run.original_query == "报销咋整"


def test_history_reaches_the_rewriter(tmp_path) -> None:
    corpus = _corpus(tmp_path)
    child_ids = _child_ids(corpus)
    rewriter = FixedRewriter("第三条的内容是什么")
    retriever = Retriever(
        "ws",
        corpus_loader=lambda: corpus,
        embeddings=FakeEmbeddings(),
        vector_store=FakeVectorStore(child_ids),
        reranker=ScriptedReranker(order=[0], scores=[0.9]),
        rewriter=rewriter,
    )
    retriever.search("那第三条呢", history=["用户: 第二条怎么规定", "助手: 第二条是..."] )
    assert rewriter.seen[0][0] == "那第三条呢"
    assert rewriter.seen[0][1] == ["用户: 第二条怎么规定", "助手: 第二条是..."]


def test_rewrite_is_skipped_when_disabled(tmp_path) -> None:
    corpus = _corpus(tmp_path)
    child_ids = _child_ids(corpus)
    store = FakeVectorStore(child_ids)
    rewriter = FixedRewriter("不该被用到")

    retriever = Retriever(
        "ws",
        corpus_loader=lambda: corpus,
        embeddings=FakeEmbeddings(),
        vector_store=store,
        reranker=ScriptedReranker(order=[0], scores=[0.9]),
        rewriter=rewriter,
    )
    retriever.search("原问题", use_rewrite=False)

    assert rewriter.seen == []
    assert store.queries == ["原问题"]


# --------------------------------------------------------------------------- #
# rewriter unit behaviour
# --------------------------------------------------------------------------- #
class _ScriptedLLM:
    def __init__(self, content) -> None:
        self.content = content
        self.calls = 0

    def invoke(self, messages):
        self.calls += 1
        if isinstance(self.content, Exception):
            raise self.content
        from langchain_core.messages import AIMessage

        return AIMessage(content=self.content)


def test_rewriter_strips_quotes_and_takes_first_line() -> None:
    rewriter = QueryRewriter(llm=_ScriptedLLM('"报销流程是什么"\n多余的解释'))
    result = rewriter.rewrite("报销咋整")
    assert result.query == "报销流程是什么"
    assert result.rewritten is True


def test_rewriter_falls_back_on_exception() -> None:
    rewriter = QueryRewriter(llm=_ScriptedLLM(RuntimeError("network down")))
    result = rewriter.rewrite("报销咋整")
    assert result.query == "报销咋整"
    assert result.rewritten is False
    assert result.error


def test_rewriter_rejects_empty_output() -> None:
    rewriter = QueryRewriter(llm=_ScriptedLLM("   "))
    result = rewriter.rewrite("报销咋整")
    assert result.query == "报销咋整"
    assert result.rewritten is False


def test_rewriter_rejects_runaway_output() -> None:
    rewriter = QueryRewriter(llm=_ScriptedLLM("很长" * 500))
    result = rewriter.rewrite("报销咋整")
    assert result.query == "报销咋整"
    assert result.rewritten is False


def test_rewriter_disabled_makes_no_model_call() -> None:
    from app.config import get_settings

    settings = get_settings()
    llm = _ScriptedLLM("不该被调用")
    rewriter = QueryRewriter(settings, llm=llm)
    original = settings.query_rewrite_enabled
    settings.query_rewrite_enabled = False
    try:
        result = rewriter.rewrite("报销咋整")
    finally:
        settings.query_rewrite_enabled = original
    assert result.query == "报销咋整"
    assert llm.calls == 0
