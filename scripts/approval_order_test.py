# -*- coding: utf-8 -*-
"""Drive a real multi-proposal chat turn and apply the proposals in a chosen order.

Usage:
    python scripts/approval_order_test.py chat <tag>          # create ws, upload, chat, save state
    python scripts/approval_order_test.py apply <tag> 0,1,2   # apply proposals in this order
    python scripts/approval_order_test.py dump <tag>          # dump the final table
"""
import json
import shutil
import sys
from pathlib import Path

import httpx

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "data" / "workspaces" / "9871b0b0d2894b71949cbf293d13e82b" / "01-产品订单表.xlsx"
STATE_DIR = ROOT / "data" / "approval_order_test"
BASE = "http://localhost:8000/api"

sys.path.insert(0, str(ROOT / "ai-service"))

client = httpx.Client(base_url=BASE, trust_env=False, timeout=600.0)


def sse_events(resp):
    event = None
    for line in resp.iter_lines():
        if not line:
            continue
        if line.startswith("event:"):
            event = line[6:].strip()
        elif line.startswith("data:"):
            payload = line[5:].strip()
            try:
                yield event, json.loads(payload)
            except json.JSONDecodeError:
                yield event, {"raw": payload}


def cmd_chat(tag: str) -> None:
    STATE_DIR.mkdir(parents=True, exist_ok=True)
    ws = client.post(
        "/workspaces", json={"name": f"审批顺序{tag}", "description": "approval-order test"}
    ).json()
    wid = ws["id"]
    with open(SRC, "rb") as fh:
        up = client.post(
            f"/workspaces/{wid}/files",
            files={"files": ("01-产品订单表.xlsx", fh, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
    up.raise_for_status()
    print("workspace", wid, "upload:", [(r["rel_path"], r["status"]) for r in up.json()])

    events = []
    with client.stream(
        "POST",
        f"/workspaces/{wid}/chat/stream",
        json={"message": "删除单价大于300的产品然后新增充电宝，单价500"},
    ) as resp:
        print("chat http", resp.status_code)
        for event, data in sse_events(resp):
            events.append({"event": event, "data": data})
            etype = event or data.get("type")
            if etype == "proposal":
                print("PROPOSAL:", json.dumps(data, ensure_ascii=False)[:600])
            elif etype == "done":
                print("DONE:", str(data.get("content"))[:300])
            elif etype == "error":
                print("ERROR:", data)

    (STATE_DIR / f"{tag}.json").write_text(
        json.dumps({"workspace_id": wid, "events": events}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print("state saved:", STATE_DIR / f"{tag}.json")


def proposals_of(tag: str) -> tuple[str, list[dict]]:
    state = json.loads((STATE_DIR / f"{tag}.json").read_text(encoding="utf-8"))
    props = [e["data"] for e in state["events"] if e["event"] == "proposal"]
    return state["workspace_id"], props


def cmd_apply(tag: str, order: str) -> None:
    wid, props = proposals_of(tag)
    print(f"{len(props)} proposals; apply order {order}")
    for idx in [int(i) for i in order.split(",")]:
        op = props[idx]
        print(f"--- applying #{idx} {op.get('tool')} op={op.get('operation_id')} summary={op.get('summary')}")
        r = client.post(f"/workspaces/{wid}/operations/{op['operation_id']}/apply")
        print("   http", r.status_code, r.json().get("status") if r.status_code == 200 else r.text[:300])


def dump_rows(path: Path):
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=False)
    for ws in wb.worksheets:
        print("SHEET:", ws.title, "dims:", ws.dimensions)
        for row in ws.iter_rows():
            vals = [(c.coordinate, c.value) for c in row if c.value is not None]
            if vals:
                print(vals)


def cmd_dump(tag: str) -> None:
    wid, _ = proposals_of(tag)
    files = client.get(f"/workspaces/{wid}/files").json()
    for f in files:
        print("== file:", f["rel_path"])
        dump_rows(ROOT / "data" / "workspaces" / wid / f["rel_path"])


if __name__ == "__main__":
    cmd = sys.argv[1]
    if cmd == "chat":
        cmd_chat(sys.argv[2])
    elif cmd == "apply":
        cmd_apply(sys.argv[2], sys.argv[3])
    elif cmd == "dump":
        cmd_dump(sys.argv[2])
    else:
        print(__doc__)
