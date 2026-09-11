"""Generate the demo documents used in the README walkthrough.

The data is deliberately seeded with rows that violate the policy in the Word
document, so the "read the policy, then fix the spreadsheet" flow has something
real to act on during a demo.

Usage:
    python scripts/make_demo_data.py [output_dir]
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

HEADER_FILL = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")


def build_workbook(path: Path) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "报销明细"

    headers = ["报销单号", "姓名", "部门", "报销类型", "金额", "提交日期", "状态"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")

    rows = [
        ("BX-2026-001", "张伟", "销售部", "差旅费", 3200, "2026-03-02", "待审批"),
        ("BX-2026-002", "李娜", "市场部", "招待费", 8600, "2026-03-04", "待审批"),
        ("BX-2026-003", "王强", "技术部", "差旅费", 1800, "2026-03-05", "已通过"),
        ("BX-2026-004", "赵敏", "销售部", "办公用品", 450, "2026-03-06", "已通过"),
        ("BX-2026-005", "陈晨", "市场部", "差旅费", 7200, "2026-03-09", "待审批"),
        ("BX-2026-006", "刘洋", "技术部", "培训费", 5600, "2026-03-11", "待审批"),
        ("BX-2026-007", "孙婷", "财务部", "办公用品", 320, "2026-03-12", "已通过"),
        ("BX-2026-008", "周杰", "销售部", "招待费", 4800, "2026-03-15", "待审批"),
    ]
    for row in rows:
        sheet.append(list(row))

    widths = [16, 10, 10, 12, 10, 14, 12]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width

    summary = book.create_sheet("部门汇总")
    summary.append(["部门", "报销笔数", "合计金额"])
    for cell in summary[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    for department in ("销售部", "市场部", "技术部", "财务部"):
        summary.append([department, None, None])
    summary["B2"] = "=COUNTIF(报销明细!C:C,A2)"
    summary["C2"] = "=SUMIF(报销明细!C:C,A2,报销明细!E:E)"
    summary["B3"] = "=COUNTIF(报销明细!C:C,A3)"
    summary["C3"] = "=SUMIF(报销明细!C:C,A3,报销明细!E:E)"
    summary["B4"] = "=COUNTIF(报销明细!C:C,A4)"
    summary["C4"] = "=SUMIF(报销明细!C:C,A4,报销明细!E:E)"
    summary["B5"] = "=COUNTIF(报销明细!C:C,A5)"
    summary["C5"] = "=SUMIF(报销明细!C:C,A5,报销明细!E:E)"
    summary.column_dimensions["A"].width = 12
    summary.column_dimensions["B"].width = 12
    summary.column_dimensions["C"].width = 16

    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)


def build_document(path: Path) -> None:
    document = Document()
    document.add_heading("员工费用报销管理制度", level=1)

    document.add_paragraph("第一条 适用范围")
    document.add_paragraph(
        "本制度适用于公司全体员工因公发生的差旅费、招待费、办公用品费及培训费用的报销。"
    )

    document.add_paragraph("第二条 报销限额")
    document.add_paragraph("单笔报销金额不得超过 5000 元，超出部分需提交总经理书面审批。")
    document.add_paragraph("其中：差旅费单笔不得超过 5000 元；招待费单笔不得超过 3000 元。")
    document.add_paragraph("办公用品费单笔不得超过 2000 元；培训费单笔不得超过 5000 元。")

    document.add_paragraph("第三条 提交时限")
    document.add_paragraph("费用发生后 30 个自然日内需提交报销申请，逾期不予受理。")

    document.add_paragraph("第四条 审批流程")
    document.add_paragraph("单笔 2000 元以下由部门负责人审批。")
    document.add_paragraph("单笔 2000 元至 5000 元由部门负责人与财务负责人共同审批。")
    document.add_paragraph("单笔 5000 元以上需经总经理审批。")

    document.add_paragraph("第五条 必备材料")
    document.add_paragraph("报销时须附发票原件、费用明细清单及对应的审批记录。")

    table = document.add_table(rows=1, cols=3)
    table.style = "Light Grid Accent 1"
    header = table.rows[0].cells
    header[0].text = "费用类型"
    header[1].text = "单笔限额（元）"
    header[2].text = "审批层级"
    for row in (
        ("差旅费", "5000", "部门负责人 + 财务负责人"),
        ("招待费", "3000", "部门负责人 + 财务负责人"),
        ("办公用品费", "2000", "部门负责人"),
        ("培训费", "5000", "部门负责人 + 财务负责人"),
    ):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value

    document.add_paragraph("第六条 附则")
    document.add_paragraph("本制度自 2026 年 1 月 1 日起施行，由财务部负责解释。")

    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))


def main() -> int:
    output = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parents[1] / "demo"
    workbook = output / "2026年第一季度报销明细.xlsx"
    doc = output / "员工费用报销管理制度.docx"
    build_workbook(workbook)
    build_document(doc)
    print(f"generated:\n  {workbook}\n  {doc}")
    print()
    print("演示流程建议：")
    print("  1. 新建工作区，把这两个文件上传")
    print("  2. 提问：制度里规定的招待费单笔上限是多少？")
    print("  3. 提问：帮我找出报销明细里超过制度限额的记录")
    print("  4. 让 Agent 把这些超标记录的状态改为「需总经理审批」，确认 Diff 后应用")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
