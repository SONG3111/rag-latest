"""Generate a manual test corpus for the workspace UI.

Files land in ``testdata/`` at the project root. Each file targets a specific
surface (MCP read/write tools, retrieval & chunking) and pairs with the prompt
list in ``testdata/README.md``.

    python scripts/make_manual_test_files.py
"""

from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "testdata"

ORDERS = [
    ("无线耳机", "数码", 12, 299.0, "2026-08-02"),
    ("机械键盘", "数码", 8, 459.0, "2026-08-03"),
    ("便携音箱", "数码", 15, 199.0, "2026-08-05"),
    ("智能手环", "数码", 20, 399.0, "2026-08-08"),
    ("降噪耳机", "数码", 6, 1299.0, "2026-08-09"),
    ("蓝牙鼠标", "数码", 25, 89.0, "2026-08-12"),
    ("USB-C 扩展坞", "配件", 18, 259.0, "2026-08-14"),
    ("显示器支架", "配件", 10, 169.0, "2026-08-15"),
    ("笔记本支架", "配件", 14, 129.0, "2026-08-18"),
    ("移动电源", "配件", 30, 149.0, "2026-08-20"),
    ("桌面麦克风", "数码", 5, 699.0, "2026-08-22"),
    ("高清摄像头", "数码", 9, 349.0, "2026-08-25"),
    ("键盘膜", "配件", 40, 19.0, "2026-08-28"),
    ("鼠标垫", "配件", 50, 29.0, "2026-08-30"),
    ("读卡器", "配件", 22, 59.0, "2026-09-01"),
    ("HDMI 线", "配件", 35, 39.0, "2026-09-02"),
]

LONG_REMARK = (
    "本批次货物按框架协议交付，具体约定如下：卖方应在收到预付款后的十五个工作日内完成首批备货，"
    "并以书面形式通知买方验收；运输方式为公路整车运输，运费由卖方承担；货到买方指定仓库后，"
    "买方应在三个工作日内完成数量与外观验收，逾期未提出书面异议视为验收合格；"
    "若发生包装破损、数量短缺或型号不符，卖方应在接到通知后的五个工作日内完成补发或换货，"
    "由此产生的额外运费由责任方承担；本批次产品的质保期为自验收合格之日起十二个月，"
    "质保期内非人为损坏的，卖方负责免费维修或更换；付款方式为预付百分之三十、货到验收合格后支付百分之六十、"
    "剩余百分之十作为质保金在质保期满后十个工作日内支付；本条与其他条款不一致的，以本条为准。"
    "关于交付进度的补充约定：首批交付不少于总量的百分之六十，第二批在首批验收合格后二十日内交付，"
    "两批之间卖方须提前五日书面告知买方具体的到货时间与车牌信息，便于买方安排收货人员与月台；"
    "若因卖方原因导致单批延迟交付超过十个工作日，买方有权按延迟天数每日收取合同金额千分之五的违约金，"
    "违约金累计不超过合同总金额的百分之十；若延迟超过三十个工作日，买方有权解除本批次订单并要求返还已付款项；"
    "关于包装与标识的约定：所有产品须采用出口级纸箱包装，箱内附防潮剂与合格证，"
    "外箱须以不褪色油墨印刷产品型号、数量、生产批次与收货仓库代码，标识不清导致的错收由卖方负责；"
    "关于售后服务的约定：卖方在买方所在城市须有常驻或合作的服务网点，接到报修后四十八小时内响应，"
    "一般故障七十二小时内修复，无法现场修复的应提供同型号备机；质保期外维修仅收取成本费，"
    "并在报价得到买方书面确认后实施；本备注未尽事宜，双方按框架协议及现行法律协商解决。"
    "关于验收标准的补充说明：外观验收以无划痕、无变形、无掉漆为合格基准，功能验收以双方确认的测试用例为准，"
    "每批次抽检比例为百分之十，抽检不合格率超过百分之二的，该批次整批判退并计入供应商季度考核；"
    "关于发票与对账的约定：卖方应在每批货款支付前开具增值税专用发票，票面信息须与订单一致，"
    "发票丢失或作废重开的，由此产生的税差由卖方承担；双方每月末对账一次，对账单经双方经办人签字盖章后生效；"
    "关于保密的约定：双方对在履约过程中知悉的对方商业信息、价格体系与客户资料负有保密义务，"
    "保密期限为合同终止后两年；违反保密义务给对方造成损失的，按实际损失承担赔偿责任；"
    "关于争议解决的约定：因本备注及框架协议引起的争议，双方应友好协商解决，"
    "协商不成的，任何一方可向买方所在地有管辖权的人民法院提起诉讼；诉讼期间，本备注不涉争议的部分应继续履行。"
)


def cell(ws, row: int, col: int, value, *, bold=False):
    c = ws.cell(row=row, column=col, value=value)
    if bold:
        c.font = Font(bold=True)
    return c


def make_order_workbook(path: Path) -> None:
    """标准订单表：单表、16 行数据、公式列、日期列。适配全部写入类工具。"""
    book = Workbook()
    ws = book.active
    ws.title = "订单"
    ws.append(["产品", "类别", "数量", "单价", "金额", "下单日期", "备注"])
    for i, (name, cat, qty, price, day) in enumerate(ORDERS, start=2):
        ws.append([name, cat, qty, price, f"=C{i}*D{i}", date.fromisoformat(day), None])
    for col in range(1, 8):
        cell(ws, 1, col, ws.cell(row=1, column=col).value, bold=True)
    ws.column_dimensions["G"].width = 40
    book.save(path)


def make_quarter_workbook(path: Path) -> None:
    """多工作表 + 跨表公式：验证 get_doc_structure 与跨 sheet 读取。"""
    book = Workbook()
    quarters = {
        "Q1": [("华东", 1200, 98.0), ("华北", 860, 102.0), ("华南", 1450, 88.5)],
        "Q2": [("华东", 1350, 95.0), ("华北", 990, 105.0), ("华南", 1520, 90.0)],
        "Q3": [("华东", 1490, 92.0), ("华北", 1080, 99.0), ("华南", 1610, 87.5)],
    }
    for name, rows in quarters.items():
        ws = book.create_sheet(name)
        ws.append(["区域", "销量", "均价"])
        for r, (region, vol, price) in enumerate(rows, start=2):
            ws.append([region, vol, price, f"=B{r}*C{r}"])
        ws.cell(row=1, column=4, value="销售额")
        for col in range(1, 5):
            ws.cell(row=1, column=col).font = Font(bold=True)

    summary = book.create_sheet("汇总", 0)
    book.remove(book["Sheet"])
    summary.append(["区域", "Q1销售额", "Q2销售额", "Q3销售额", "前三季度合计"])
    regions = ["华东", "华北", "华南"]
    for r, region in enumerate(regions, start=2):
        summary.append([
            region,
            f"=Q1!D{r}",
            f"=Q2!D{r}",
            f"=Q3!D{r}",
            f"=SUM(B{r}:D{r})",
        ])
    summary.append(["合计", "=SUM(B2:B4)", "=SUM(C2:C4)", "=SUM(D2:D4)", "=SUM(E2:E4)"])
    for col in range(1, 6):
        summary.cell(row=1, column=col).font = Font(bold=True)
    book.save(path)


def make_adversarial_workbook(path: Path) -> None:
    """刁钻结构：同表堆叠双表、偏移表头、超长单元格、宽表。验证分块修复。"""
    book = Workbook()

    stacked = book.active
    stacked.title = "两个表"
    stacked.append(["产品", "类别", "数量", "单价"])
    stacked.append(["显示器", "数码", 7, 899.0])
    stacked.append(["打印机", "办公", 3, 1299.0])
    stacked.append(["碎纸机", "办公", 5, 459.0])
    stacked.append([None, None, None, None])
    stacked.append(["部门", "报销人", "事由", "金额", "状态"])
    stacked.append(["市场部", "张伟", "客户招待", 1200.0, "已审批"])
    stacked.append(["技术部", "李娜", "差旅住宿", 860.0, "待审批"])
    stacked.append(["人事部", "王强", "培训费", 2400.0, "已审批"])
    stacked.append(["财务部", "赵敏", "办公用品", 320.0, "已驳回"])

    offset = book.create_sheet("偏移数据")
    offset["B3"] = "姓名"
    offset["C3"] = "城市"
    offset["D3"] = "报销类型"
    offset["B4"] = "陈晨"
    offset["C4"] = "杭州"
    offset["D4"] = "差旅费"
    offset["B5"] = "林一"
    offset["C5"] = "宁波"
    offset["D5"] = "招待费"

    long_text = book.create_sheet("长单元格")
    long_text.append(["单号", "备注"])
    long_text.append(["PO-2026-001", LONG_REMARK])

    wide = book.create_sheet("宽表")
    wide.append([f"指标{chr(65 + i // 26)}{chr(65 + i % 26) if i >= 26 else ''}" for i in range(30)])
    for row_index in range(1, 6):
        wide.append([f"值{row_index}-{j}" for j in range(30)])
    book.save(path)


def make_format_workbook(path: Path) -> None:
    """留给 format_range 练手的素表：无任何格式。"""
    book = Workbook()
    ws = book.active
    ws.title = "考核"
    ws.append(["员工", "部门", "得分", "达标线"])
    rows = [
        ("陈晨", "市场部", 88, 80),
        ("李娜", "技术部", 92, 85),
        ("王强", "人事部", 76, 80),
        ("赵敏", "财务部", 95, 90),
        ("孙磊", "技术部", 81, 85),
    ]
    for r in rows:
        ws.append(list(r))
    book.save(path)


def make_policy_doc(path: Path) -> None:
    """制度文档：章节 + 数字条款 + 表格。适配检索问答与 replace_text。"""
    doc = Document()
    doc.add_heading("公司报销管理制度", level=1)

    doc.add_heading("第一章 总则", level=2)
    doc.add_paragraph("本制度适用于公司全体员工及劳务派遣人员，自发布之日起施行。")
    doc.add_paragraph("报销货币为人民币，以外币结算的费用按报销当月一日的汇率折算。")

    doc.add_heading("第二章 报销标准", level=2)
    doc.add_paragraph("单笔报销金额不得超过 5000 元，超出部分需总经理审批。")
    doc.add_paragraph("出差补贴标准为每日 200 元，按实际出差天数计算，往返当日各按半天计。")
    doc.add_paragraph(
        "招待费单笔上限为 2000 元，陪餐人数不得超过我方出差人数的三倍；"
        "超过 2000 元的招待费须事先经部门负责人书面同意。"
    )
    doc.add_paragraph("市场一部人员的差旅报销由区域经理复核后统一提交财务部。")

    doc.add_heading("第三章 审批流程", level=2)
    doc.add_paragraph("报销单据须在费用发生后三十个工作日内提交，逾期财务部门有权拒绝受理。")
    doc.add_paragraph("审批链为：部门负责人 → 财务部 → 总经理（仅限超限单据）。")

    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "职级"
    table.rows[0].cells[1].text = "住宿标准（元/晚）"
    table.rows[0].cells[2].text = "交通标准（元/公里）"
    for level in range(1, 11):
        row = table.add_row().cells
        row[0].text = f"P{level}"
        row[1].text = str(300 + level * 50)
        row[2].text = str(1 + level // 10)
    doc.save(path)


def make_plan_doc(path: Path) -> None:
    """计划书：多级标题 + 超长段落（触发 token 切分）+ 自定义标题样式章节。"""
    doc = Document()
    doc.add_heading("智慧仓储项目计划书", level=1)

    doc.add_heading("一、项目背景", level=2)
    doc.add_paragraph("现有仓库依赖人工盘点，账实一致率长期在 92% 左右，旺季峰值拣货错误率上升明显。")
    doc.add_paragraph("本项目目标是在 2026 年第四季度前将账实一致率提升到 99.5%，拣货效率提升 30%。")

    doc.add_heading("二、实施方案", level=2)
    long_paragraph = (
        "实施方案分为三个阶段推进。第一阶段为基建与设备进场，预计四周，包括货架改造、网络布线、"
        "基站部署与叉车改装；设备到货后由供应商现场指导安装，安装完成的每个库区须经过三天的空载联调，"
        "确认标签读取率不低于 99.9% 后方可进入带载测试；带载测试期间将抽取真实订单的 10% 进行灰度跑单，"
        "灰度期间拣货错误率若高于现状则暂停放量并回溯问题；第二阶段为系统对接与流程切换，预计六周，"
        "包括仓储管理系统与现有进销存系统的接口开发、库存盘点流程重造、员工培训与持证上岗考核；"
        "接口开发采用先读后写的方式灰度上线，写接口须经过双方业务负责人确认的字段映射表方可投产；"
        "第三阶段为全量切换与优化，预计四周，切换选择在月末盘点后的低峰期进行，"
        "切换完成后保留旧流程双轨运行两周以应对异常回退；双轨期内每日对账一次，"
        "差异率连续五天低于千分之三方可拆除旧流程；项目全程设立由仓储部、信息部与供应商组成的三方例会，"
        "每周五对进度、风险与预算执行情况进行书面复盘，重大变更须经项目管理委员会审批后执行。"
    )
    doc.add_paragraph(long_paragraph)

    doc.add_heading("三、预算与里程碑", level=2)
    doc.add_paragraph("项目总预算 186 万元，其中设备 120 万、软件 40 万、实施服务 26 万。")
    doc.add_paragraph("关键里程碑：9 月 30 日完成基建，11 月 15 日完成系统对接，12 月 20 日全量切换。")

    # 最后一章用自定义标题样式（基于 Heading 1），验证本地化/自定义标题识别
    style = doc.styles.add_style("附则标题", WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = doc.styles["Heading 1"]
    doc.add_paragraph("四、附则", style=style)
    doc.add_paragraph("本计划书由项目管理委员会负责解释，未尽事宜另行通知。")
    doc.save(path)


def make_minutes_doc(path: Path) -> None:
    """会议纪要：两张表格。适配 read_table / update_table_cell。"""
    doc = Document()
    doc.add_heading("仓储升级项目第 3 次周会纪要", level=1)
    doc.add_paragraph("时间：2026 年 9 月 8 日 14:00-15:20；地点：三号会议室。")

    doc.add_heading("参会人员", level=2)
    t1 = doc.add_table(rows=1, cols=3)
    t1.style = "Table Grid"
    t1.rows[0].cells[0].text = "姓名"
    t1.rows[0].cells[1].text = "部门"
    t1.rows[0].cells[2].text = "角色"
    for name, dept, role in [
        ("陈晨", "仓储部", "项目经理"),
        ("李娜", "信息部", "系统对接负责人"),
        ("王强", "仓储部", "设备负责人"),
    ]:
        row = t1.add_row().cells
        row[0].text, row[1].text, row[2].text = name, dept, role

    doc.add_heading("行动项", level=2)
    t2 = doc.add_table(rows=1, cols=4)
    t2.style = "Table Grid"
    for i, header in enumerate(["编号", "事项", "负责人", "截止日期"]):
        t2.rows[0].cells[i].text = header
    for no, task, owner, due in [
        ("1", "完成库区 B 的基站联调", "王强", "2026-09-15"),
        ("2", "提交进销存接口字段映射表", "李娜", "2026-09-18"),
        ("3", "灰度跑单错误率周报", "陈晨", "2026-09-12"),
    ]:
        row = t2.add_row().cells
        row[0].text, row[1].text, row[2].text, row[3].text = no, task, owner, due

    doc.add_heading("决议", level=2)
    doc.add_paragraph("灰度跑单比例维持 10%，9 月底前不再上调。")
    doc.save(path)


def main() -> int:
    OUT.mkdir(exist_ok=True)
    make_order_workbook(OUT / "01-产品订单表.xlsx")
    make_quarter_workbook(OUT / "02-季度销售汇总.xlsx")
    make_adversarial_workbook(OUT / "03-复杂结构表.xlsx")
    make_format_workbook(OUT / "04-考核评分表.xlsx")
    make_policy_doc(OUT / "05-公司报销管理制度.docx")
    make_plan_doc(OUT / "06-智慧仓储项目计划书.docx")
    make_minutes_doc(OUT / "07-仓储升级周会纪要.docx")
    for f in sorted(OUT.iterdir()):
        print(f"{f.name:<28} {f.stat().st_size:>7} B")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
