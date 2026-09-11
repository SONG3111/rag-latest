from __future__ import annotations

import sys
from pathlib import Path

import pytest

SRC = Path(__file__).resolve().parents[1] / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))


@pytest.fixture()
def workspace(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    """An isolated workspace root honoured by the sandbox."""
    root = tmp_path / "workspace"
    root.mkdir()
    monkeypatch.setenv("WORKSPACE_ROOT", str(root))
    return root


@pytest.fixture()
def workbook_path(workspace: Path) -> Path:
    from openpyxl import Workbook

    wb = Workbook()
    sheet = wb.active
    sheet.title = "销售"
    sheet.append(["产品", "区域", "销售额"])
    sheet.append(["A型", "华东", 1000])
    sheet.append(["B型", "华北", 2000])
    sheet.append(["C型", "华南", 3000])
    summary = wb.create_sheet("汇总")
    summary["A1"] = "合计"
    summary["B1"] = "=SUM(销售!C2:C4)"
    path = workspace / "销售表.xlsx"
    wb.save(path)
    return path


@pytest.fixture()
def word_path(workspace: Path) -> Path:
    from docx import Document

    document = Document()
    document.add_heading("报销制度", level=1)
    document.add_paragraph("第一条 员工出差需提前申请。")
    document.add_paragraph("第二条 单笔报销金额不得超过 5000 元。")
    document.add_heading("附则", level=2)
    document.add_paragraph("本制度自发布之日起执行。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "限额"
    table.cell(0, 1).text = "金额"
    table.cell(1, 0).text = "单笔报销"
    table.cell(1, 1).text = "5000"
    path = workspace / "制度.docx"
    document.save(str(path))
    return path
