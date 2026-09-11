from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from mcp_office_server import word_ops
from mcp_office_server.errors import ParagraphNotFound, RangeError, TableNotFound


def test_structure_reports_outline_and_tables(word_path: Path) -> None:
    structure = word_ops.document_structure(word_path)
    assert structure["kind"] == "word"
    headings = [item["text"] for item in structure["outline"]]
    assert headings == ["报销制度", "附则"]
    assert structure["outline"][0]["level"] == 1
    assert structure["outline"][1]["level"] == 2
    assert structure["table_count"] == 1
    assert structure["tables"][0]["header"] == ["限额", "金额"]


def test_read_paragraphs_slices_by_index(word_path: Path) -> None:
    result = word_ops.read_paragraphs(word_path, start=0, end=2)
    assert result["returned"] == 3
    assert result["paragraphs"][0]["text"] == "报销制度"
    assert result["paragraphs"][0]["is_heading"] is True
    assert result["paragraphs"][1]["is_heading"] is False


def test_read_paragraphs_rejects_out_of_range_start(word_path: Path) -> None:
    with pytest.raises(ParagraphNotFound):
        word_ops.read_paragraphs(word_path, start=999)


def test_read_table_returns_grid(word_path: Path) -> None:
    result = word_ops.read_table(word_path, 0)
    assert result["rows"] == [["限额", "金额"], ["单笔报销", "5000"]]


def test_read_table_rejects_unknown_index(word_path: Path) -> None:
    with pytest.raises(TableNotFound):
        word_ops.read_table(word_path, 7)


def test_find_text_spans_paragraphs_and_tables(word_path: Path) -> None:
    hits = word_ops.find_text(word_path, "5000")
    locations = {(hit["location"]) for hit in hits}
    assert locations == {"paragraph", "table_cell"}


def test_replace_text_updates_body_and_table(word_path: Path) -> None:
    result = word_ops.replace_text(word_path, "5000", "8000")
    assert result["total_replacements"] == 2

    document = Document(str(word_path))
    texts = [paragraph.text for paragraph in document.paragraphs]
    assert any("8000" in text for text in texts)
    assert document.tables[0].cell(1, 1).text == "8000"


def test_replace_text_matches_across_run_boundaries(word_path: Path) -> None:
    """Word splits a sentence into several runs; replacement must still hit."""
    document = Document(str(word_path))
    paragraph = document.paragraphs[1]
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)
    for piece in ("第一条 ", "员工出差", "需提前申请。"):
        paragraph.add_run(piece)
    document.save(str(word_path))
    assert len(Document(str(word_path)).paragraphs[1].runs) == 3

    result = word_ops.replace_text(word_path, "员工出差需提前申请", "员工出差须提前审批")
    assert result["total_replacements"] == 1

    updated = Document(str(word_path)).paragraphs[1].text
    assert updated == "第一条 员工出差须提前审批。"


def test_replace_text_reports_nothing_when_absent(word_path: Path) -> None:
    result = word_ops.replace_text(word_path, "不存在的词", "x")
    assert result["total_replacements"] == 0
    assert result["changes"] == []


def test_replace_text_rejects_empty_find(word_path: Path) -> None:
    with pytest.raises(RangeError):
        word_ops.replace_text(word_path, "", "x")


def test_update_table_cell_reports_diff(word_path: Path) -> None:
    result = word_ops.update_table_cell(word_path, 0, 1, 1, "6000")
    change = result["changes"][0]
    assert change["before"] == "5000"
    assert change["after"] == "6000"
    assert Document(str(word_path)).tables[0].cell(1, 1).text == "6000"


def test_update_table_cell_rejects_out_of_range(word_path: Path) -> None:
    with pytest.raises(RangeError):
        word_ops.update_table_cell(word_path, 0, 9, 0, "x")
    with pytest.raises(TableNotFound):
        word_ops.update_table_cell(word_path, 3, 0, 0, "x")
