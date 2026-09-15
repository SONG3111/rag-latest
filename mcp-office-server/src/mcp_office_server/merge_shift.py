"""Shift merged cell ranges after row/column insert/delete operations.

openpyxl moves cell *values* on ``insert_rows``/``delete_rows`` (and the column
variants) but leaves merged ranges untouched — the same documented limitation
as formula strings (openpyxl docs, "Inserting and deleting rows and columns":
"Openpyxl does not manage dependencies, such as formulae, tables, charts, …";
tracked as openpyxl issue #1139). Excel itself keeps merged headers aligned
with their data, so after a shift on this server the merge would otherwise sit
over the wrong rows — e.g. deleting a merged header row leaves the merge
covering whatever data moved up into it.

The rules replicate Excel's own behaviour, reusing the same band arithmetic as
``formula_shift`` so both layers stay symmetric:

* insert ``count`` lines at ``start`` — a merge at/after ``start`` moves
  forward by ``count``; a merge straddling ``start`` (the band opens inside
  it) is *extended* by ``count`` (inserting a row inside a merged header grows
  the merge, as in Excel); a merge entirely above is untouched.
* delete lines ``[start, start+count)`` — a merge entirely inside the band is
  *removed* together with its rows; a merge entirely below moves back by
  ``count``; a merge partially covered shrinks: the edge inside the band
  collapses onto the band's near edge (low edge to ``start``, high edge to
  ``start-1``), the outer edge moves back by ``count`` — the same edge
  collapse ``formula_shift`` applies to range references.

The unmerge → move → re-merge pattern is the standard community workaround for
this openpyxl limitation (see e.g. haris-musa/excel-mcp-server, which shares
the limitation, and StackOverflow "Openpyxl Issue: Deleting row not moving
merged cells", question 53162099).

Known leftovers, same scope note as ``formula_shift``: charts, tables,
defined names and images are not rewritten on shifts.
"""

from __future__ import annotations

from typing import Literal

from openpyxl.worksheet.cell_range import CellRange
from openpyxl.worksheet.worksheet import Worksheet as OpenpyxlWorksheet

ShiftOp = Literal["insert", "delete"]
ShiftAxis = Literal["row", "col"]


def _band(r: CellRange, axis: ShiftAxis) -> tuple[int, int]:
    """The merge's (low, high) coordinate along the shifted axis."""
    if axis == "row":
        return r.min_row, r.max_row
    return r.min_col, r.max_col


def _moved_bounds(r: CellRange, *, op: ShiftOp, start: int, count: int,
                  axis: ShiftAxis) -> tuple[int, int, int, int] | None:
    """New (min_col, min_row, max_col, max_row) for the merge, or None to drop it.

    Coordinates are the merge's *original* bounds: openpyxl has already moved
    the values but never the merges, so the stored ranges are still in their
    pre-shift positions when this runs.
    """
    lo, hi = _band(r, axis)

    if op == "insert":
        if lo >= start:
            new_lo, new_hi = lo + count, hi + count
        elif hi >= start:  # band opens inside the merge: Excel extends it
            new_lo, new_hi = lo, hi + count
        else:
            return r.bounds
    else:
        band_end = start + count
        if lo >= start and hi < band_end:
            return None  # entirely inside the deleted band
        if lo >= band_end:
            # entirely below the band: the whole merge moves back by count
            new_lo, new_hi = lo - count, hi - count
        elif lo >= start:
            # low edge inside, high edge beyond: the low edge collapses onto
            # the band's first line, the content below moves up under it
            new_lo, new_hi = start, hi - count
        elif hi >= band_end:
            # merge spans the whole band: shrink from the high side
            new_lo, new_hi = lo, hi - count
        elif hi >= start:
            # high edge inside: collapse onto the band's top
            new_lo, new_hi = lo, start - 1
        else:
            return r.bounds  # entirely above the band

    if axis == "row":
        return (r.min_col, new_lo, r.max_col, new_hi)
    return (new_lo, r.min_row, new_hi, r.max_row)


def shift_merged_ranges(sheet: OpenpyxlWorksheet, *, op: ShiftOp,
                        start: int, count: int, axis: ShiftAxis = "row") -> dict:
    """Apply the merge-range shift for one structural operation, in place.

    Returns small counters so callers can surface what happened:
    ``{"shifted": n, "removed": n}`` — ``shifted`` counts merges whose
    coordinates changed (moved, extended or shrunk), ``removed`` merges
    deleted together with their rows/columns.
    """
    shifted = removed = 0
    # Snapshot first: the manipulation mutates the collection being iterated.
    # Ranges are detached directly instead of ``unmerge_cells``: the structural
    # operation has already moved/deleted the underlying cells, so the strict
    # cell cleanup inside unmerge would KeyError on the stale coordinates.
    # ``merge_cells`` then re-anchors the merge over the moved cells.
    for r in list(sheet.merged_cells.ranges):
        target = _moved_bounds(r, op=op, start=start, count=count, axis=axis)
        if target == r.bounds:
            continue
        sheet.merged_cells.ranges.remove(r)
        if target is None:
            removed += 1
        else:
            new_range = CellRange(
                min_col=target[0], min_row=target[1],
                max_col=target[2], max_row=target[3],
            )
            sheet.merge_cells(new_range.coord)
            shifted += 1
    return {"shifted": shifted, "removed": removed}


__all__ = ["shift_merged_ranges"]
