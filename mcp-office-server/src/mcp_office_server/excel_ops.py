"""Excel implementation layer built on openpyxl.

This module intentionally contains no MCP tool definitions, so it can be unit
tested and reused independently of the protocol layer. Tool wrappers in
``server.py`` are thin adapters over these functions.
"""

from __future__ import annotations

import hashlib
import re
from copy import copy
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.formula.translate import Translator
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.worksheet.worksheet import Worksheet as OpenpyxlWorksheet

from . import formula_shift
from . import merge_shift
from .errors import (
    CorruptDocument,
    FileLocked,
    InvalidSheetName,
    InvalidValue,
    LastSheetError,
    MergeConflict,
    RangeError,
    SheetExists,
    SheetNotFound,
    ToolError,
    WriteConflict,
)

MAX_PREVIEW_ROWS = 200
MAX_PREVIEW_COLS = 60
MAX_MERGED_PREVIEW = 20
MAX_REPLACE_HITS = 200
MAX_FORMULA_REFS = 50

# Excel's own sheet dimensions (the same numbers openpyxl enforces when a cell
# is materialised, see ``Worksheet._get_cell``). Validated up-front so an
# out-of-limits request fails with an actionable error instead of a raw
# ValueError or, worse, a file Excel refuses to open.
EXCEL_MAX_ROW = 1048576
EXCEL_MAX_COLUMN = 16384  # column XFD

# Excel specification limits ("Excel specifications and limits"): a cell holds
# at most 32,767 characters and a formula is at most 8,192 characters long.
# openpyxl does not enforce either — worse, it *silently truncates* longer
# cell strings at save time, so the tool would report a value the file does
# not contain. Both are validated up-front and rejected instead.
MAX_CELL_CHARS = 32_767
MAX_FORMULA_CHARS = 8_192

# The control characters openpyxl itself refuses to store
# (``ILLEGAL_CHARACTERS_RE`` above); tab/newline/carriage return are legal.
# Checked here so the agent gets an actionable error naming the offending
# input instead of a raw IllegalCharacterError traceback.
_ILLEGAL_CHARS = ILLEGAL_CHARACTERS_RE

# Formatting is per-cell work in openpyxl; a whole-sheet request would spin for
# minutes, so cap the area per call the way read_range caps its preview.
MAX_FORMAT_CELLS = 200_000

# Mirrored from openpyxl: the Alignment horizontal value set and the sheet-title
# character blacklist (``openpyxl.workbook.child.INVALID_TITLE_REGEX``), plus
# the aRGB colour shape openpyxl itself enforces at assignment time.
_HEX_COLOR = re.compile(r"^(?:[0-9A-Fa-f]{6}|[0-9A-Fa-f]{8})$")
_HORIZONTAL_ALIGNMENTS = frozenset(
    {"general", "left", "center", "right", "fill", "justify",
     "centerContinuous", "distributed"}
)
_INVALID_SHEET_TITLE = re.compile(r"[\\*?:/\[\]]")


# --------------------------------------------------------------------------- #
# helpers
# --------------------------------------------------------------------------- #
def file_digest(path: Path) -> str:
    """SHA-256 of the file contents, used for optimistic write-conflict checks."""
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(65536), b""):
            digest.update(chunk)
    return digest.hexdigest()


def open_workbook(path: Path, *, read_only: bool = False, data_only: bool = False):
    """Open a workbook, converting openpyxl failures into our error type."""
    try:
        return load_workbook(path, read_only=read_only, data_only=data_only)
    except Exception as exc:  # openpyxl raises a wide variety of errors
        raise CorruptDocument(f"cannot open workbook '{path.name}': {exc}") from exc


def save_workbook(workbook, path: Path) -> None:
    """Persist the workbook, translating OS-level write failures into ToolError.

    On Windows an openpyxl save raises ``PermissionError`` when the user has the
    file open in Excel/WPS. Left uncaught, FastMCP wraps it into a plain-text
    result that the caller cannot distinguish from success.
    """
    try:
        workbook.save(path)
    except PermissionError as exc:
        raise FileLocked(
            "文件正被其他程序占用（如 Excel/WPS），请关闭后重试",
            detail=str(exc),
        ) from exc
    except OSError as exc:
        raise ToolError(f"写入文件失败: {exc}", detail=str(exc)) from exc


def _require_sheet(workbook, sheet_name: str) -> OpenpyxlWorksheet:
    if sheet_name not in workbook.sheetnames:
        raise SheetNotFound(
            f"sheet '{sheet_name}' not found; available: {', '.join(workbook.sheetnames)}"
        )
    return workbook[sheet_name]


def _parse_range(range_text: str) -> tuple[int, int, int, int]:
    """Return normalized ``(min_col, min_row, max_col, max_row)`` for an A1 range.

    openpyxl's ``range_boundaries`` returns endpoints in the order written, so a
    reversed range like ``C4:A2`` yields negative widths; Excel treats reversed
    selections as their normalized rectangle, so we sort the endpoints. Bare
    column/row spans (``A:B``) come back with ``None`` components and row 0
    parses fine — both are rejected as actionably invalid.
    """
    try:
        start_col, start_row, end_col, end_row = range_boundaries(range_text)
    except Exception as exc:
        raise RangeError(f"invalid range '{range_text}'") from exc
    if None in (start_col, start_row, end_col, end_row):
        raise RangeError(
            f"invalid range '{range_text}': use explicit cell coordinates "
            f"(e.g. 'A1:C4'), not whole-column/row spans like 'A:B'"
        )
    if min(start_row, end_row) < 1 or min(start_col, end_col) < 1:
        raise RangeError(f"invalid range '{range_text}': rows and columns start at 1")
    return (
        min(start_col, end_col),
        min(start_row, end_row),
        max(start_col, end_col),
        max(start_row, end_row),
    )


def _check_bounds(bounds: tuple[int, int, int, int], what: str) -> None:
    """Reject ranges that exceed Excel's sheet dimensions before any mutation."""
    _, _, max_col, max_row = bounds
    if max_row > EXCEL_MAX_ROW or max_col > EXCEL_MAX_COLUMN:
        raise RangeError(
            f"{what} exceeds Excel's sheet limits "
            f"(max row {EXCEL_MAX_ROW}, max column XFD={EXCEL_MAX_COLUMN})"
        )


def _jsonify(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def _validate_text(value: str) -> str:
    """Reject text Excel cannot hold verbatim, so the file matches the report.

    openpyxl silently truncates strings longer than Excel's 32,767-character
    cell limit at save time (the tool would report the full text while the
    file keeps a cut-off copy) and raises a raw ``IllegalCharacterError`` for
    control characters. Both are turned into actionable ``invalid_value``
    errors before anything is written; tab/newline/carriage return stay legal.
    """
    if len(value) > MAX_CELL_CHARS:
        raise InvalidValue(
            f"text is {len(value)} characters; an Excel cell holds at most "
            f"{MAX_CELL_CHARS} — split it across cells or shorten it"
        )
    illegal = _ILLEGAL_CHARS.search(value)
    if illegal:
        raise InvalidValue(
            f"text contains the control character {illegal.group(0)!r} "
            f"(U+{ord(illegal.group(0)):04X}), which Excel cannot store; "
            "remove it (\\n and \\t are fine)"
        )
    return value


def _coerce_input(value: Any) -> Any:
    """Normalize values coming from the model before writing them to a cell.

    Models occasionally wrap a formula in a small object (``{"formula": "=A1*2"}``)
    instead of passing the string; the intent is unambiguous, so unwrap it rather
    than failing the whole write with openpyxl's ``Cannot convert ... to Excel``.
    Anything else an Excel cell cannot hold (nested objects, lists) is rejected
    with an actionable message instead of a raw ``ValueError``.
    """
    if value is None:
        return None
    if isinstance(value, str):
        # Only a truly empty string means "clear"; a whitespace-only string is
        # data like any other and is stored verbatim, as Excel would.
        if value == "":
            return None
        return _validate_text(value)
    if isinstance(value, dict) and set(value) == {"formula"} and isinstance(value["formula"], str):
        return value["formula"]
    if isinstance(value, (bool, int, float, datetime, date)):
        return value
    raise InvalidValue(
        f"a cell can hold text, a number, a boolean, a date or a formula string; "
        f"got {type(value).__name__}"
    )


def _validate_sheet_name(name: str) -> None:
    """Reject sheet names Excel itself would refuse to open or round-trip.

    Character blacklist and the 31-character limit are Excel's own rules,
    mirrored from ``openpyxl.workbook.child`` (which only *warns* on long
    titles — Excel hard-rejects them, so we do too).
    """
    if not name or not name.strip():
        raise InvalidSheetName("sheet name must not be empty")
    if len(name) > 31:
        raise InvalidSheetName(
            f"sheet name {name[:16]}… is {len(name)} characters; Excel limits names to 31"
        )
    match = _INVALID_SHEET_TITLE.search(name)
    if match:
        raise InvalidSheetName(
            f"sheet name cannot contain {match.group(0)!r}; "
            "forbidden characters are \\ * ? : / [ ]"
        )


def _check_conflict(path: Path, expected_digest: str | None) -> None:
    if not expected_digest:
        return
    actual = file_digest(path)
    if actual != expected_digest:
        raise WriteConflict(
            "the file changed on disk since it was read; re-read it before applying changes"
        )


def _parse_cell(cell_text: str) -> tuple[int, int]:
    """Return ``(col, row)`` for a single A1-style coordinate."""
    min_col, min_row, max_col, max_row = _parse_range(cell_text)
    if (min_col, min_row) != (max_col, max_row):
        raise RangeError(f"'{cell_text}' is a range; a single cell is required here")
    return min_col, min_row


def _merged_overlapping(
    sheet: OpenpyxlWorksheet, bounds: tuple[int, int, int, int]
) -> CellRange | None:
    """First merged range overlapping ``(min_col, min_row, max_col, max_row)``, if any."""
    min_col, min_row, max_col, max_row = bounds
    for merged in sheet.merged_cells.ranges:
        if (
            min_col <= merged.max_col
            and max_col >= merged.min_col
            and min_row <= merged.max_row
            and max_row >= merged.min_row
        ):
            return merged
    return None


def _reject_non_anchor_merge_write(
    sheet: OpenpyxlWorksheet, col: int, row: int, coordinate: str
) -> None:
    """Refuse writes into a merged region except its top-left anchor.

    Excel keeps only the anchor writable; openpyxl would raise a raw
    AttributeError on the MergedCell otherwise. Messages name the range and
    its anchor so the agent can correct the coordinate in one step.
    """
    merged = _merged_overlapping(sheet, (col, row, col, row))
    if merged is None:
        return
    if (col, row) != (merged.min_col, merged.min_row):
        anchor = f"{get_column_letter(merged.min_col)}{merged.min_row}"
        raise MergeConflict(
            f"cell '{coordinate}' is inside the merged range {merged}; only its "
            f"top-left cell {anchor} is writable — unmerge first to change the others"
        )


def _require_no_merge(sheet: OpenpyxlWorksheet, bounds: tuple[int, int, int, int], what: str) -> None:
    conflict = _merged_overlapping(sheet, bounds)
    if conflict:
        raise MergeConflict(
            f"{what} crosses the merged range {conflict}; unmerge it first "
            f"(unmerge_cells) or use update_cells for individual values"
        )


def _formulas_referencing(workbook, sheet_title: str) -> list[str]:
    """Formulas anywhere in the workbook that reference ``sheet_title``.

    Used as a warning when deleting/renaming a sheet: openpyxl does not rewrite
    cross-sheet references, so Excel will show #REF! for these afterwards.
    """
    needle = sheet_title.casefold()
    hits: list[str] = []
    for ws in workbook.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if not (isinstance(value, str) and value.startswith("=")):
                    continue
                folded = value.casefold()
                if f"{needle}!" in folded or f"'{needle}'!" in folded:
                    hits.append(f"{ws.title}!{cell.coordinate}")
                    if len(hits) >= MAX_FORMULA_REFS:
                        return hits
    return hits


# --------------------------------------------------------------------------- #
# read operations
# --------------------------------------------------------------------------- #
def workbook_structure(path: Path) -> dict:
    """Describe every sheet: dimensions, header row, merged ranges, widths.

    Opened in normal (not read-only) mode because ``merged_cells`` is not
    available there — merged ranges carry table semantics (大类/小类 headers)
    the agent needs before writing.
    """
    workbook = open_workbook(path)
    try:
        sheets = []
        for name in workbook.sheetnames:
            sheet = workbook[name]
            header: list[Any] = []
            for row in sheet.iter_rows(min_row=1, max_row=1, max_col=min(sheet.max_column or 0, MAX_PREVIEW_COLS)):
                header = [_jsonify(cell.value) for cell in row]
                break
            merged = [str(merged) for merged in sheet.merged_cells.ranges]
            sheets.append(
                {
                    "name": name,
                    "max_row": sheet.max_row or 0,
                    "max_column": sheet.max_column or 0,
                    "header": header,
                    "merged_ranges": merged[:MAX_MERGED_PREVIEW],
                    "merged_truncated": len(merged) > MAX_MERGED_PREVIEW,
                }
            )
        return {"kind": "excel", "sheets": sheets}
    finally:
        workbook.close()


def read_range(
    path: Path,
    sheet_name: str,
    *,
    start_cell: str = "A1",
    end_cell: str | None = None,
    max_rows: int = MAX_PREVIEW_ROWS,
    max_cols: int = MAX_PREVIEW_COLS,
) -> dict:
    """Read a rectangular region, returning raw values plus formula visibility.

    Values are read with ``data_only=False`` so formulas are visible to the agent
    as formula text; a cached value is reported alongside when available.
    """
    workbook = open_workbook(path, data_only=False)
    cached = open_workbook(path, data_only=True)
    try:
        sheet = _require_sheet(workbook, sheet_name)
        cached_sheet = cached[sheet_name] if sheet_name in cached.sheetnames else None

        # Semantics ported from haris-musa/excel-mcp-server
        # (``read_excel_range_with_metadata``): a range may be passed as
        # ``start_cell="A1:C4"``, and an omitted ``end_cell`` means "expand to the
        # sheet's data", not "read one cell".
        if ":" in start_cell and not end_cell:
            start_cell, end_cell = (part.strip() for part in start_cell.split(":", 1))

        if end_cell:
            min_col, min_row, max_col, max_row = _parse_range(f"{start_cell}:{end_cell}")
        else:
            min_col, min_row, _, _ = _parse_range(start_cell)
            if start_cell.upper() == "A1":
                # A sheet whose data does not start at A1 should still read whole.
                min_row, min_col = sheet.min_row, sheet.min_column
            empty = (
                (sheet.max_row or 0) <= 1
                and (sheet.max_column or 0) <= 1
                and sheet.cell(row=1, column=1).value is None
            )
            if empty:
                max_row, max_col = min_row, min_col
            else:
                max_row = max(sheet.max_row or min_row, min_row)
                max_col = max(sheet.max_column or min_col, min_col)
        _check_bounds((min_col, min_row, max_col, max_row), f"range {start_cell}:{end_cell or start_cell}")

        total_rows = (max_row - min_row + 1)
        total_cols = (max_col - min_col + 1)
        if total_rows <= 0 or total_cols <= 0:
            raise RangeError(f"empty range {start_cell}:{end_cell or start_cell}")

        truncated = total_rows > max_rows or total_cols > max_cols
        last_row = min(max_row, min_row + max_rows - 1)
        last_col = min(max_col, min_col + max_cols - 1)

        rows: list[list[Any]] = []
        for row_idx in range(min_row, last_row + 1):
            values: list[Any] = []
            for col_idx in range(min_col, last_col + 1):
                cell = sheet.cell(row=row_idx, column=col_idx)
                raw = cell.value
                if isinstance(raw, str) and raw.startswith("="):
                    cached_value = (
                        cached_sheet.cell(row=row_idx, column=col_idx).value
                        if cached_sheet is not None
                        else None
                    )
                    values.append(
                        {"formula": raw, "cached_value": _jsonify(cached_value)}
                    )
                else:
                    values.append(_jsonify(raw))
            rows.append(values)

        return {
            "kind": "excel",
            "sheet": sheet_name,
            # The effective range, not the request: when the caller passes A1 on a
            # sheet whose data starts at B3, the answer should say so explicitly.
            "requested_start_cell": start_cell,
            "start_cell": f"{get_column_letter(min_col)}{min_row}",
            "end_cell": f"{get_column_letter(last_col)}{last_row}",
            "values": rows,
            "truncated": truncated,
            "total_rows": total_rows,
            "total_columns": total_cols,
            "sheet_max_row": sheet.max_row or 0,
            "sheet_max_column": sheet.max_column or 0,
            "digest": file_digest(path),
        }
    finally:
        workbook.close()
        cached.close()


def find_text(path: Path, query: str, *, max_hits: int = 50) -> list[dict]:
    """Case-insensitive substring search across all sheets.

    Non-text cells are matched against their string form — the idiom openpyxl
    based servers use upstream — because a user asking "哪里出现了 299" expects
    the cell holding 299.0 to be found, not a silent miss. Dates match their
    ISO form and formulas match their formula text.
    """
    needle = query.strip().casefold()
    if not needle:
        return []
    hits: list[dict] = []
    workbook = open_workbook(path, read_only=True)
    try:
        for name in workbook.sheetnames:
            sheet = workbook[name]
            for row in sheet.iter_rows():
                for cell in row:
                    value = cell.value
                    if value is None:
                        continue
                    if needle in str(_jsonify(value)).casefold():
                        hits.append(
                            {
                                "sheet": name,
                                "cell": cell.coordinate,
                                "value": _jsonify(value),
                            }
                        )
                        if len(hits) >= max_hits:
                            return hits
        return hits
    finally:
        workbook.close()


# --------------------------------------------------------------------------- #
# write operations
# --------------------------------------------------------------------------- #
def update_cells(
    path: Path,
    sheet_name: str,
    updates: Iterable[dict],
    *,
    expected_digest: str | None = None,
) -> dict:
    """Set cell values, returning a per-cell before/after diff.

    Each update is ``{"cell": "B3", "value": ...}``. ``None`` or ``""`` clears
    a cell; whitespace-only text is data and is stored verbatim, like Excel.
    """
    _check_conflict(path, expected_digest)
    workbook = open_workbook(path)
    try:
        sheet = _require_sheet(workbook, sheet_name)
        changes: list[dict] = []
        for update in updates:
            coordinate = str(update.get("cell", "")).strip()
            if not coordinate:
                raise RangeError("each update requires a 'cell' coordinate")
            col, row = _parse_cell(coordinate)
            _check_bounds((col, row, col, row), f"cell '{coordinate}'")

            target = sheet.cell(row=row, column=col)
            _reject_non_anchor_merge_write(sheet, col, row, coordinate)
            before = _jsonify(target.value)
            after = _coerce_input(update.get("value"))
            target.value = after
            changes.append(
                {
                    "cell": coordinate,
                    "before": before,
                    "after": _jsonify(after),
                }
            )
        save_workbook(workbook, path)
        return {
            "kind": "excel",
            "sheet": sheet_name,
            "path": path.name,
            "changes": changes,
            "digest": file_digest(path),
        }
    finally:
        workbook.close()


def set_formula(
    path: Path,
    sheet_name: str,
    cell: str,
    formula: str,
    *,
    expected_digest: str | None = None,
) -> dict:
    """Write an Excel formula, normalizing the leading ``=``."""
    expression = formula.strip()
    if not expression:
        raise RangeError("formula must not be empty")
    body = expression[1:] if expression.startswith("=") else expression
    if not body.strip():
        # openpyxl stores a bare "=" as literal text, so the tool would report a
        # formula where the file has a string — reject it the way Excel does.
        raise RangeError(
            f"formula {formula!r} has no body after '='; "
            "write the expression, e.g. 'SUM(B2:B10)'"
        )
    if len(expression) > MAX_FORMULA_CHARS:
        raise RangeError(
            f"formula is {len(expression)} characters; Excel limits formulas "
            f"to {MAX_FORMULA_CHARS}"
        )
    illegal = _ILLEGAL_CHARS.search(expression)
    if illegal:
        raise RangeError(
            f"formula contains the control character {illegal.group(0)!r}, "
            "which Excel cannot store"
        )
    if not expression.startswith("="):
        expression = f"={expression}"

    _check_conflict(path, expected_digest)
    workbook = open_workbook(path)
    try:
        sheet = _require_sheet(workbook, sheet_name)
        col, row = _parse_cell(cell)
        _check_bounds((col, row, col, row), f"cell '{cell}'")
        target = sheet.cell(row=row, column=col)
        _reject_non_anchor_merge_write(sheet, col, row, cell)
        before = _jsonify(target.value)
        target.value = expression
        save_workbook(workbook, path)
        return {
            "kind": "excel",
            "sheet": sheet_name,
            "changes": [{"cell": cell, "before": before, "after": expression}],
            "digest": file_digest(path),
        }
    finally:
        workbook.close()


def insert_rows(
    path: Path,
    sheet_name: str,
    start_row: int,
    count: int = 1,
    *,
    expected_digest: str | None = None,
) -> dict:
    if start_row < 1:
        raise RangeError("start_row is 1-based and must be >= 1")
    if count < 1:
        raise RangeError("count must be >= 1")
    if start_row > EXCEL_MAX_ROW:
        raise RangeError(f"start_row {start_row} is beyond Excel's last row ({EXCEL_MAX_ROW})")

    _check_conflict(path, expected_digest)
    workbook = open_workbook(path)
    try:
        sheet = _require_sheet(workbook, sheet_name)
        # Pushing data past Excel's last row silently produces a file Excel
        # refuses to open; bound-checking the result keeps the file valid.
        if (sheet.max_row or 0) + count > EXCEL_MAX_ROW:
            raise RangeError(
                f"inserting {count} rows would push the sheet past Excel's "
                f"{EXCEL_MAX_ROW}-row limit; insert fewer rows"
            )
        before_rows = sheet.max_row or 0
        sheet.insert_rows(start_row, amount=count)
        # openpyxl moves cells but not formula strings; rewrite references the
        # way Excel does (see formula_shift).
        formula_shift.shift_workbook_formulas(
            workbook, affected_sheet=sheet.title, op="insert",
            start=start_row, count=count,
        )
        # ...nor merged ranges; move/extend them like Excel does (merge_shift).
        merged = merge_shift.shift_merged_ranges(
            sheet, op="insert", start=start_row, count=count, axis="row",
        )
        save_workbook(workbook, path)
        return {
            "kind": "excel",
            "sheet": sheet_name,
            "operation": "insert_rows",
            "start_row": start_row,
            "count": count,
            "rows_before": before_rows,
            "rows_after": sheet.max_row or 0,
            "merged_shifted": merged["shifted"],
            "merged_removed": merged["removed"],
            "digest": file_digest(path),
        }
    finally:
        workbook.close()


def delete_rows(
    path: Path,
    sheet_name: str,
    start_row: int,
    count: int = 1,
    *,
    expected_digest: str | None = None,
) -> dict:
    if start_row < 1:
        raise RangeError("start_row is 1-based and must be >= 1")
    if count < 1:
        raise RangeError("count must be >= 1")

    _check_conflict(path, expected_digest)
    workbook = open_workbook(path)
    try:
        sheet = _require_sheet(workbook, sheet_name)
        before_rows = sheet.max_row or 0
        sheet.delete_rows(start_row, amount=count)
        formula_shift.shift_workbook_formulas(
            workbook, affected_sheet=sheet.title, op="delete",
            start=start_row, count=count,
        )
        merged = merge_shift.shift_merged_ranges(
            sheet, op="delete", start=start_row, count=count, axis="row",
        )
        save_workbook(workbook, path)
        return {
            "kind": "excel",
            "sheet": sheet_name,
            "operation": "delete_rows",
            "start_row": start_row,
            "count": count,
            "rows_before": before_rows,
            "rows_after": sheet.max_row or 0,
            "merged_shifted": merged["shifted"],
            "merged_removed": merged["removed"],
            "digest": file_digest(path),
        }
    finally:
        workbook.close()


def format_range(
    path: Path,
    sheet_name: str,
    start_cell: str,
    end_cell: str | None = None,
    *,
    bold: bool | None = None,
    italic: bool | None = None,
    font_color: str | None = None,
    fill_color: str | None = None,
    number_format: str | None = None,
    horizontal_alignment: str | None = None,
    expected_digest: str | None = None,
) -> dict:
    """Apply basic cell formatting to a rectangle."""
    for label, color in (("font_color", font_color), ("fill_color", fill_color)):
        if color is not None and not _HEX_COLOR.match(color):
            raise RangeError(
                f"{label} must be RGB hex like 'FF0000' (optionally 8-digit aRGB); got {color!r}"
            )
    if horizontal_alignment is not None and horizontal_alignment not in _HORIZONTAL_ALIGNMENTS:
        raise RangeError(
            f"horizontal_alignment must be one of {', '.join(sorted(_HORIZONTAL_ALIGNMENTS))}; "
            f"got {horizontal_alignment!r}"
        )

    _check_conflict(path, expected_digest)
    workbook = open_workbook(path)
    try:
        sheet = _require_sheet(workbook, sheet_name)
        min_col, min_row, max_col, max_row = _parse_range(
            f"{start_cell}:{end_cell}" if end_cell else start_cell
        )
        _check_bounds((min_col, min_row, max_col, max_row), f"range {start_cell}:{end_cell or start_cell}")
        if (max_row - min_row + 1) * (max_col - min_col + 1) > MAX_FORMAT_CELLS:
            # Per-cell styling over billions of coordinates would spin for
            # minutes and exhaust memory; bounded the same way reads are.
            raise RangeError(
                f"range {start_cell}:{end_cell or start_cell} covers more than "
                f"{MAX_FORMAT_CELLS} cells; narrow the range or format it in batches"
            )

        applied: list[str] = []
        for row in sheet.iter_rows(
            min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col
        ):
            for cell in row:
                if bold is not None or italic is not None or font_color is not None:
                    font = cell.font
                    cell.font = Font(
                        name=font.name,
                        size=font.size,
                        bold=bold if bold is not None else font.bold,
                        italic=italic if italic is not None else font.italic,
                        color=font_color or font.color,
                    )
                if fill_color:
                    cell.fill = PatternFill(
                        start_color=fill_color, end_color=fill_color, fill_type="solid"
                    )
                if number_format:
                    cell.number_format = number_format
                if horizontal_alignment:
                    cell.alignment = Alignment(horizontal=horizontal_alignment)

        if bold is not None:
            applied.append(f"bold={bold}")
        if italic is not None:
            applied.append(f"italic={italic}")
        if font_color:
            applied.append(f"font_color={font_color}")
        if fill_color:
            applied.append(f"fill_color={fill_color}")
        if number_format:
            applied.append(f"number_format={number_format}")
        if horizontal_alignment:
            applied.append(f"alignment={horizontal_alignment}")

        save_workbook(workbook, path)
        return {
            "kind": "excel",
            "sheet": sheet_name,
            "range": f"{start_cell}:{end_cell or start_cell}",
            "applied": applied,
            "digest": file_digest(path),
        }
    finally:
        workbook.close()


# --------------------------------------------------------------------------- #
# structural write operations (column / range / merge / sheet level)
# --------------------------------------------------------------------------- #
def insert_columns(
    path: Path,
    sheet_name: str,
    start_col: int,
    count: int = 1,
    *,
    expected_digest: str | None = None,
) -> dict:
    """Insert empty columns before ``start_col``, rewriting formula references.

    The column twin of ``insert_rows``. Use this to add columns to a table;
    appending data to the right of existing data does not need it.
    """
    if start_col < 1:
        raise RangeError("start_col is 1-based and must be >= 1")
    if count < 1:
        raise RangeError("count must be >= 1")
    if start_col > EXCEL_MAX_COLUMN:
        raise RangeError(f"start_col {start_col} is beyond Excel's last column ({EXCEL_MAX_COLUMN})")

    _check_conflict(path, expected_digest)
    workbook = open_workbook(path)
    try:
        sheet = _require_sheet(workbook, sheet_name)
        if (sheet.max_column or 0) + count > EXCEL_MAX_COLUMN:
            raise RangeError(
                f"inserting {count} columns would push the sheet past Excel's "
                f"{EXCEL_MAX_COLUMN}-column limit; insert fewer columns"
            )
        before_cols = sheet.max_column or 0
        sheet.insert_cols(start_col, amount=count)
        formula_shift.shift_workbook_formulas(
            workbook, affected_sheet=sheet.title, op="insert",
            start=start_col, count=count, axis="col",
        )
        merged = merge_shift.shift_merged_ranges(
            sheet, op="insert", start=start_col, count=count, axis="col",
        )
        save_workbook(workbook, path)
        return {
            "kind": "excel",
            "sheet": sheet_name,
            "operation": "insert_columns",
            "start_col": start_col,
            "count": count,
            "columns_before": before_cols,
            "columns_after": sheet.max_column or 0,
            "merged_shifted": merged["shifted"],
            "merged_removed": merged["removed"],
            "digest": file_digest(path),
        }
    finally:
        workbook.close()


def delete_columns(
    path: Path,
    sheet_name: str,
    start_col: int,
    count: int = 1,
    *,
    expected_digest: str | None = None,
) -> dict:
    """Delete whole columns; formulas referencing them become #REF!.

    The column twin of ``delete_rows``. To blank values without shifting the
    table, use ``update_cells`` with ``null`` values instead.
    """
    if start_col < 1:
        raise RangeError("start_col is 1-based and must be >= 1")
    if count < 1:
        raise RangeError("count must be >= 1")

    _check_conflict(path, expected_digest)
    workbook = open_workbook(path)
    try:
        sheet = _require_sheet(workbook, sheet_name)
        before_cols = sheet.max_column or 0
        sheet.delete_cols(start_col, amount=count)
        formula_shift.shift_workbook_formulas(
            workbook, affected_sheet=sheet.title, op="delete",
            start=start_col, count=count, axis="col",
        )
        merged = merge_shift.shift_merged_ranges(
            sheet, op="delete", start=start_col, count=count, axis="col",
        )
        save_workbook(workbook, path)
        return {
            "kind": "excel",
            "sheet": sheet_name,
            "operation": "delete_columns",
            "start_col": start_col,
            "count": count,
            "columns_before": before_cols,
            "columns_after": sheet.max_column or 0,
            "merged_shifted": merged["shifted"],
            "merged_removed": merged["removed"],
            "digest": file_digest(path),
        }
    finally:
        workbook.close()


def copy_range(
    path: Path,
    src_sheet: str,
    src_range: str,
    dst_sheet: str,
    dst_cell: str,
    *,
    expected_digest: str | None = None,
) -> dict:
    """Copy a rectangular block (values, styles, formulas) to a destination anchor.

    Concept and semantics follow ``copy_range`` from haris-musa/excel-mcp-server
    (MIT); formulas are an improvement over the upstream verbatim copy: relative
    references are translated to the destination the way Excel does (via
    openpyxl's ``Translator``), absolute ones stay put. References without a
    sheet qualifier keep pointing at *their own* sheet inside the formula — when
    copied across sheets that becomes the destination sheet, matching Excel.

    Use ``update_cells`` for point edits of known cells; use this to duplicate
    a whole block. Ranges touching merged cells are rejected — unmerge first.
    """
    src_bounds = _parse_range(src_range)
    dst_col, dst_row = _parse_cell(dst_cell)
    src_min_col, src_min_row, src_max_col, src_max_row = src_bounds
    rows = src_max_row - src_min_row + 1
    cols = src_max_col - src_min_col + 1
    dst_bounds = (dst_col, dst_row, dst_col + cols - 1, dst_row + rows - 1)
    _check_bounds(src_bounds, f"source range '{src_range}'")
    _check_bounds(dst_bounds, f"destination area starting at '{dst_cell}'")

    _check_conflict(path, expected_digest)
    workbook = open_workbook(path)
    try:
        source = _require_sheet(workbook, src_sheet)
        target = _require_sheet(workbook, dst_sheet)
        _require_no_merge(source, src_bounds, f"source range '{src_range}'")
        _require_no_merge(target, dst_bounds, f"destination area starting at '{dst_cell}'")

        # Snapshot the whole source block before writing anything: when source
        # and destination overlap, writing while iterating would re-read cells
        # the copy itself has already overwritten (Excel snapshots first — the
        # upstream haris-musa/excel-mcp-server copy loop does not, and corrupts
        # overlapping copies).
        staged: list[tuple[bool, Any, Any, str]] = []
        for row_offset in range(rows):
            for col_offset in range(cols):
                src = source.cell(row=src_min_row + row_offset, column=src_min_col + col_offset)
                is_formula = isinstance(src.value, str) and src.value.startswith("=")
                staged.append(
                    (is_formula, src.value, copy(src._style) if src.has_style else None, src.coordinate)
                )

        for row_offset in range(rows):
            for col_offset in range(cols):
                is_formula, value, style, src_coordinate = staged[row_offset * cols + col_offset]
                dst = target.cell(row=dst_row + row_offset, column=dst_col + col_offset)
                if is_formula:
                    value = Translator(
                        value, origin=src_coordinate
                    ).translate_formula(dst.coordinate)
                dst.value = value
                if style is not None:
                    dst._style = style

        save_workbook(workbook, path)
        return {
            "kind": "excel",
            # The normalized rectangle actually copied (a reversed request like
            # C4:A2 reports A2:C4 — the same rectangle, canonical form).
            "source": (
                f"{src_sheet}!{get_column_letter(src_min_col)}{src_min_row}"
                f":{get_column_letter(src_max_col)}{src_max_row}"
            ),
            "destination": (
                f"{dst_sheet}!{get_column_letter(dst_bounds[0])}{dst_bounds[1]}"
                f":{get_column_letter(dst_bounds[2])}{dst_bounds[3]}"
            ),
            "copied_cells": rows * cols,
            "digest": file_digest(path),
        }
    finally:
        workbook.close()


def delete_range(
    path: Path,
    sheet_name: str,
    range_text: str,
    shift: str = "up",
    *,
    expected_digest: str | None = None,
) -> dict:
    """Delete a rectangular block and pull neighbours in to fill the hole.

    ``shift="up"`` moves everything below the block up; ``shift="left"`` moves
    everything right of it left. Formula references inside the moved cells are
    translated (openpyxl ``move_range`` semantics); formulas *pointing at* the
    moved cells from elsewhere are not rewritten — same limitation Excel-based
    servers inherit from openpyxl. For whole-row/column deletion use
    ``delete_rows``/``delete_columns``; to blank values without moving anything,
    use ``update_cells`` with ``null``.

    Block deletion with neighbour shift is not native to openpyxl; this
    implementation moves values, styles and formulas per column/row, mirroring
    ``delete_range(shift=...)`` from haris-musa/excel-mcp-server (MIT).
    Styles travel with their cells; leftover cells keep their old formatting.
    """
    if shift not in ("up", "left"):
        raise RangeError("shift must be 'up' or 'left'")

    min_col, min_row, max_col, max_row = _parse_range(range_text)
    _check_bounds((min_col, min_row, max_col, max_row), f"range '{range_text}'")
    _check_conflict(path, expected_digest)
    workbook = open_workbook(path)
    try:
        sheet = _require_sheet(workbook, sheet_name)
        vacated: tuple[int, int, int, int] | None = None
        if shift == "up":
            # Clamp the requested band to the sheet's used rows: a range whose
            # bottom reaches past the data would otherwise clear rows *above*
            # the band (and walk negative row indices) when the tail is emptied.
            effective_max_row = min(max_row, sheet.max_row or max_row)
            if effective_max_row < min_row:
                moved = 0  # the range sits entirely below the data: nothing to do
            else:
                tail_end = sheet.max_row or max_row
                _require_no_merge(sheet, (min_col, min_row, max_col, max(tail_end, max_row)),
                                  f"range '{range_text}' (with the cells below it)")
                moved = _shift_block_vertical(
                    sheet, min_col, max_col, min_row, effective_max_row, tail_end
                )
                # The vacated lines are the bottom ``rows`` lines of the old
                # tail (see _shift_block_vertical) — or the band itself when
                # nothing sat below it to move in. Either way they end blank,
                # so report them regardless of how many cells moved.
                rows = effective_max_row - min_row + 1
                vacated = (min_col, tail_end - rows + 1, max_col, tail_end)
        else:
            effective_max_col = min(max_col, sheet.max_column or max_col)
            if effective_max_col < min_col:
                moved = 0
            else:
                tail_end = sheet.max_column or max_col
                _require_no_merge(sheet, (min_col, min_row, max(tail_end, max_col), max_row),
                                  f"range '{range_text}' (with the cells to its right)")
                moved = _shift_block_horizontal(
                    sheet, min_row, max_row, min_col, effective_max_col, tail_end
                )
                cols = effective_max_col - min_col + 1
                vacated = (tail_end - cols + 1, min_row, tail_end, max_row)

        def _range_string(bounds: tuple[int, int, int, int] | None) -> str | None:
            if bounds is None:
                return None
            b0, b1, b2, b3 = bounds
            return (
                f"{get_column_letter(b0)}{b1}:{get_column_letter(b2)}{b3}"
            )

        save_workbook(workbook, path)
        return {
            "kind": "excel",
            "sheet": sheet_name,
            "range": range_text,
            "shift": shift,
            "moved_cells": moved,
            # Canonical rectangle of the cells left blank by the shift, e.g.
            # "A4:B5"; null when nothing was vacated.
            "vacated_range": _range_string(vacated),
            "digest": file_digest(path),
        }
    finally:
        workbook.close()


def _shift_block_vertical(sheet: OpenpyxlWorksheet, min_col: int, max_col: int,
                          min_row: int, max_row: int, tail_end: int) -> int:
    """Move rows below ``[min_row, max_row]`` up into the band, then clear the tail.

    After the move the vacated lines are the bottom ``rows`` lines of the old
    tail — not the whole tail, whose upper part now holds shifted content.
    """
    rows = max_row - min_row + 1
    moved = 0
    for col in range(min_col, max_col + 1):
        staged: list[tuple[Any, Any]] = []
        for row in range(max_row + 1, tail_end + 1):
            cell = sheet.cell(row=row, column=col)
            staged.append((cell.value, copy(cell._style) if cell.has_style else None))
        for offset, (value, style) in enumerate(staged):
            dst = sheet.cell(row=min_row + offset, column=col)
            if isinstance(value, str) and value.startswith("="):
                value = Translator(
                    value, origin=sheet.cell(row=max_row + 1 + offset, column=col).coordinate
                ).translate_formula(dst.coordinate)
            dst.value = value
            if style is not None:
                dst._style = style
            moved += 1
        for row in range(tail_end - rows + 1, tail_end + 1):
            sheet.cell(row=row, column=col).value = None
    return moved


def _shift_block_horizontal(sheet: OpenpyxlWorksheet, min_row: int, max_row: int,
                            min_col: int, max_col: int, tail_end: int) -> int:
    """Move columns right of ``[min_col, max_col]`` left into the band, then clear the tail."""
    cols = max_col - min_col + 1
    moved = 0
    for row in range(min_row, max_row + 1):
        staged: list[tuple[Any, Any]] = []
        for col in range(max_col + 1, tail_end + 1):
            cell = sheet.cell(row=row, column=col)
            staged.append((cell.value, copy(cell._style) if cell.has_style else None))
        for offset, (value, style) in enumerate(staged):
            dst = sheet.cell(row=row, column=min_col + offset)
            if isinstance(value, str) and value.startswith("="):
                value = Translator(
                    value, origin=sheet.cell(row=row, column=max_col + 1 + offset).coordinate
                ).translate_formula(dst.coordinate)
            dst.value = value
            if style is not None:
                dst._style = style
            moved += 1
        for col in range(tail_end - cols + 1, tail_end + 1):
            sheet.cell(row=row, column=col).value = None
    return moved


def merge_cells(
    path: Path,
    sheet_name: str,
    range_text: str,
    *,
    expected_digest: str | None = None,
) -> dict:
    """Merge a range into one cell; values outside the top-left are discarded.

    Same behaviour as merging in Excel: only the top-left cell keeps its value.
    Re-merging an already merged range is a no-op; overlapping a *different*
    merge is rejected (``merge_conflict``) — unmerge first.
    """
    bounds = _parse_range(range_text)
    min_col, min_row, max_col, max_row = bounds
    if min_col == max_col and min_row == max_row:
        raise RangeError(
            f"'{range_text}' is a single cell; merging needs a range of at least two cells"
        )
    _check_bounds(bounds, f"range '{range_text}'")
    _check_conflict(path, expected_digest)
    workbook = open_workbook(path)
    try:
        sheet = _require_sheet(workbook, sheet_name)
        existing = {tuple(r.bounds): str(r) for r in sheet.merged_cells.ranges}
        if bounds in existing:
            return {
                "kind": "excel",
                "sheet": sheet_name,
                "merged": sorted(str(r) for r in sheet.merged_cells.ranges),
                "already_merged": existing[bounds],
                "digest": file_digest(path),
            }
        conflict = _merged_overlapping(sheet, bounds)
        if conflict:
            raise MergeConflict(
                f"'{range_text}' overlaps the existing merge {conflict}; unmerge it first"
            )
        sheet.merge_cells(range_text)
        save_workbook(workbook, path)
        return {
            "kind": "excel",
            "sheet": sheet_name,
            "merged": sorted(str(r) for r in sheet.merged_cells.ranges),
            "digest": file_digest(path),
        }
    finally:
        workbook.close()


def unmerge_cells(
    path: Path,
    sheet_name: str,
    range_text: str,
    *,
    expected_digest: str | None = None,
) -> dict:
    """Unmerge a merged range; the range must match an existing merge exactly.

    Like Excel, a partial selection cannot be unmerged: pass the full merged
    range (they are listed in ``get_doc_structure`` as ``merged_ranges``).
    """
    bounds = _parse_range(range_text)
    _check_conflict(path, expected_digest)
    workbook = open_workbook(path)
    try:
        sheet = _require_sheet(workbook, sheet_name)
        existing = {tuple(r.bounds): str(r) for r in sheet.merged_cells.ranges}
        if bounds not in existing:
            listing = ", ".join(sorted(existing.values())) or "none"
            raise RangeError(
                f"no merged range exactly matches '{range_text}'; existing merges: {listing}"
            )
        sheet.unmerge_cells(existing[bounds])
        save_workbook(workbook, path)
        return {
            "kind": "excel",
            "sheet": sheet_name,
            "unmerged": existing[bounds],
            "merged": sorted(str(r) for r in sheet.merged_cells.ranges),
            "digest": file_digest(path),
        }
    finally:
        workbook.close()


def find_replace(
    path: Path,
    query: str,
    replacement: str,
    *,
    sheet_name: str | None = None,
    include_formulas: bool = False,
    max_replacements: int = MAX_REPLACE_HITS,
    expected_digest: str | None = None,
) -> dict:
    """Replace a substring across text cells, returning every touched cell.

    Matching is case-insensitive and mirrors ``find_text`` so what the agent
    found is what this replaces. Only string cells are touched — numbers and
    dates keep their type; formula strings are only rewritten when
    ``include_formulas`` is true (their cached values refresh on next open).
    Use ``update_cells`` when you already know the exact cells; use this for
    bulk renames like a product or department name change.
    """
    query = query.strip()
    if not query:
        raise RangeError("query must not be empty")
    if not isinstance(replacement, str):
        raise RangeError("replacement must be a string")

    pattern = re.compile(re.escape(query), re.IGNORECASE)

    def literal_replace(_match: re.Match) -> str:
        # The replacement is data, not a regex template: pass it through a
        # callable so backslashes (\1, \g<...>, or a Windows path's trailing
        # backslash) are inserted literally instead of raising "bad escape" or
        # expanding groups — the standard idiom for literal substitution with
        # re.sub (re.escape must not be used for the replacement side).
        return replacement
    _check_conflict(path, expected_digest)
    workbook = open_workbook(path)
    try:
        if sheet_name is not None:
            targets = [_require_sheet(workbook, sheet_name)]
        else:
            targets = list(workbook.worksheets)
        changes: list[dict] = []
        truncated = False
        for sheet in targets:
            for row in sheet.iter_rows():
                for cell in row:
                    value = cell.value
                    if not isinstance(value, str) or not value:
                        continue
                    if value.startswith("=") and not include_formulas:
                        continue
                    new_value, hits = pattern.subn(literal_replace, value)
                    if not hits:
                        continue
                    changes.append(
                        {
                            "sheet": sheet.title,
                            "cell": cell.coordinate,
                            "before": value,
                            "after": new_value,
                            "occurrences": hits,
                        }
                    )
                    cell.value = new_value
                    if len(changes) >= max_replacements:
                        truncated = True
                        break
                if truncated:
                    break
            if truncated:
                break

        if changes:
            save_workbook(workbook, path)
        return {
            "kind": "excel",
            "query": query,
            "replacement": replacement,
            "changes": changes,
            "total_replacements": sum(c["occurrences"] for c in changes),
            "truncated": truncated,
            "digest": file_digest(path),
        }
    finally:
        workbook.close()


def manage_sheets(
    path: Path,
    action: str,
    sheet_name: str | None = None,
    *,
    new_name: str | None = None,
    expected_digest: str | None = None,
) -> dict:
    """Create, rename, copy or delete a worksheet in place.

    ``action`` selects the operation (the multi-mode pattern ``calculate``
    uses): ``create`` needs ``new_name``; ``rename`` and ``copy`` need both
    ``sheet_name`` and ``new_name``; ``delete`` needs ``sheet_name``.
    openpyxl does not rewrite cross-sheet references, so rename/delete return
    the formulas still pointing at the old name as a warning list.
    """
    if action not in ("create", "rename", "copy", "delete"):
        raise RangeError(f"unknown action '{action}'; expected create|rename|copy|delete")

    _check_conflict(path, expected_digest)
    workbook = open_workbook(path)
    try:
        stale_refs: list[str] = []
        if action == "create":
            if not new_name:
                raise RangeError("create requires 'new_name'")
            _validate_sheet_name(new_name)
            if new_name in workbook.sheetnames:
                raise SheetExists(f"sheet '{new_name}' already exists")
            workbook.create_sheet(title=new_name)
            target = new_name
        elif action == "rename":
            sheet = _require_sheet(workbook, sheet_name or "")
            if not new_name:
                raise RangeError("rename requires 'new_name'")
            _validate_sheet_name(new_name)
            if new_name != sheet.title and new_name in workbook.sheetnames:
                raise SheetExists(f"sheet '{new_name}' already exists")
            stale_refs = _formulas_referencing(workbook, sheet.title) if new_name != sheet.title else []
            old_title = sheet.title
            sheet.title = new_name
            target = old_title
        elif action == "copy":
            sheet = _require_sheet(workbook, sheet_name or "")
            if not new_name:
                raise RangeError("copy requires 'new_name' for the clone")
            _validate_sheet_name(new_name)
            if new_name in workbook.sheetnames:
                raise SheetExists(f"sheet '{new_name}' already exists")
            clone = workbook.copy_worksheet(sheet)
            clone.title = new_name
            target = new_name
        else:  # delete
            sheet = _require_sheet(workbook, sheet_name or "")
            if len(workbook.sheetnames) <= 1:
                raise LastSheetError("cannot delete the only sheet in the workbook")
            stale_refs = _formulas_referencing(workbook, sheet.title)
            workbook.remove(sheet)
            target = sheet.title

        save_workbook(workbook, path)
        result: dict[str, Any] = {
            "kind": "excel",
            "action": action,
            "sheet": target,
            "sheets": workbook.sheetnames,
            "digest": file_digest(path),
        }
        if stale_refs:
            result["stale_formula_refs"] = stale_refs
            result["warning"] = (
                f"openpyxl does not rewrite cross-sheet references; "
                f"{len(stale_refs)} formula(s) still point at '{target}' and will show #REF!"
            )
        return result
    finally:
        workbook.close()


__all__ = [
    "copy_range",
    "delete_columns",
    "delete_range",
    "delete_rows",
    "file_digest",
    "find_replace",
    "find_text",
    "format_range",
    "insert_columns",
    "insert_rows",
    "manage_sheets",
    "merge_cells",
    "read_range",
    "set_formula",
    "unmerge_cells",
    "update_cells",
    "workbook_structure",
]
